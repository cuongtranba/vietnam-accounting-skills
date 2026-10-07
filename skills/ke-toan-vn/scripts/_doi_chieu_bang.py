"""Động cơ `doi-chieu`: đối chiếu hai nguồn theo khoá, theo một trong hai chế độ.

`gop`       cộng cột giá trị theo khoá ở mỗi bên rồi so tổng (vd: số lượng chế phẩm máu theo
            TENVP giữa file viện phí và sổ xuất nhập tồn). Nếu khai báo `chi_tiet` (vd MABN),
            các cặp (khoá, chi tiết) lệch được liệt kê ngay cạnh dòng lệch để lần ra nguyên nhân.
`tung_dong` ghép từng dòng A với một dòng B chưa dùng cùng khoá — mỗi dòng B chỉ được dùng
            một lần, để bệnh nhân có hai lượt cùng dịch vụ không bị tính khớp hai lần vào cùng
            một dòng (vd: so bản lọc của script với bản kế toán tự làm tay).

Mỗi bên có thể loại bớt dòng bằng `loai_tru` (cùng cú pháp quy tắc với động cơ `loc`).
Khác `doi_chieu.py` (bảng kê hoá đơn ↔ sổ): không bỏ số 0 đầu của khoá trừ khi được yêu cầu,
vì mã khoa/mã bệnh nhân có số 0 đầu là có nghĩa.
"""

from __future__ import annotations

from pathlib import Path

from _quy_trinh import (
    KHOA_CHUNG, LoiQuyTrinh, bam, chi_so_cot, chuan, chuan_quy_tac, cot_cua_quy_tac,
    doc_bang, khop_quy_tac, kiem_khoa, la_so, ten_ra, ten_sheet_hop_le, tra_cot,
)

KHOA_BEN = {"nhan", "tieu_de", "sheet", "dong_tieu_de", "khoa", "chi_tiet", "gia_tri",
            "bo_so_0_dau", "loai_tru"}
KHOA_KET_QUA = {"che_do", "ra", "sheet", "thu_tu", "nguong", "cot_ghi_chu", "chi_ghi_lech"}
GIOI_HAN_CAP = 20


def kiem(qt: dict, danh_muc: dict) -> dict:
    kiem_khoa(qt, KHOA_CHUNG | {"a", "b", "ket_qua"}, qt["ten"])
    ben = {}
    for t in ("a", "b"):
        b = qt.get(t)
        if not isinstance(b, dict):
            raise LoiQuyTrinh(f"Quy trình doi-chieu cần bảng [{t}]")
        kiem_khoa(b, KHOA_BEN, f"[{t}]")
        for can in ("khoa", "gia_tri"):
            if can not in b:
                raise LoiQuyTrinh(f"[{t}] cần `{can}`")
        ben[t] = {
            "nhan": b.get("nhan", t.upper()), "sheet": b.get("sheet"),
            "tieu_de": b.get("tieu_de"),
            "dong_tieu_de": b.get("dong_tieu_de", 1),
            "khoa": list(b["khoa"]), "chi_tiet": list(b.get("chi_tiet", [])),
            "gia_tri": b["gia_tri"], "bo_so_0_dau": {chuan(c) for c in b.get("bo_so_0_dau", [])},
            "loai_tru": [chuan_quy_tac(q, danh_muc, f"[[{t}.loai_tru]] thứ {i}")
                         for i, q in enumerate(b.get("loai_tru", []), 1)],
        }
    a, b = ben["a"], ben["b"]
    if len(a["khoa"]) != len(b["khoa"]) or len(a["chi_tiet"]) != len(b["chi_tiet"]):
        raise LoiQuyTrinh("[a] và [b] phải có cùng số cột `khoa` và `chi_tiet` (ghép theo thứ tự)")
    kq = qt.get("ket_qua", {})
    kiem_khoa(kq, KHOA_KET_QUA, "[ket_qua]")
    che_do = kq.get("che_do", "gop")
    if che_do not in ("gop", "tung_dong"):
        raise LoiQuyTrinh("[ket_qua] che_do phải là \"gop\" hoặc \"tung_dong\"")
    if kq.get("thu_tu") and len(a["khoa"]) != 1:
        raise LoiQuyTrinh("`thu_tu` chỉ dùng được khi khoá có đúng một cột")
    return {**qt, "a": a, "b": b, "ket_qua": {
        "che_do": che_do, "ra": "DOI-CHIEU-{ten_file}.xlsx", "sheet": "Đối chiếu",
        "thu_tu": [], "nguong": 0.5, "cot_ghi_chu": "Ghi chú",
        "chi_ghi_lech": che_do == "tung_dong", **kq}}


def _so(v) -> float:
    if la_so(v):
        return float(v)
    if v is None or v == "":
        return 0.0
    from chuan_hoa import so_vn

    d = so_vn(v)
    return float(d) if d is not None else 0.0


def _doc_ben(ben: dict, p: Path) -> dict:
    ten_sheet, tieu_de, dong = doc_bang(p, ben["sheet"], ben["dong_tieu_de"])
    cm = chi_so_cot(tieu_de)
    can = set(ben["khoa"]) | set(ben["chi_tiet"]) | {ben["gia_tri"]}
    for q in ben["loai_tru"]:
        can |= cot_cua_quy_tac(q)
    thieu = [c for c in sorted(can) if not c.startswith("$") and chuan(c) not in cm]
    if thieu:
        raise LoiQuyTrinh(f"{p.name}/{ten_sheet}: thiếu cột {thieu}")
    vt = {c: tra_cot(cm, c, p.name) for c in can}

    def khoa_cua(r, cot):
        ra = []
        for c in cot:
            s = chuan(r[vt[c]])
            if chuan(c) in ben["bo_so_0_dau"]:
                s = s.lstrip("0") or s
            ra.append(s)
        return tuple(ra)

    giu, bo, ly_do = [], 0, {}
    for so, r in dong:
        def lay(c, r=r):
            return r[vt[c]]
        trung = [q["ma"] for q in ben["loai_tru"] if khop_quy_tac(q, lay)]
        if trung:
            bo += 1
            for m in trung:
                ly_do[m] = ly_do.get(m, 0) + 1
            continue
        k = khoa_cua(r, ben["khoa"])
        if all(x == "" for x in k):
            bo += 1
            ly_do["khoá trống"] = ly_do.get("khoá trống", 0) + 1
            continue
        giu.append({
            "so_dong": so, "khoa": k, "chi_tiet": khoa_cua(r, ben["chi_tiet"]),
            "nhan_khoa": tuple(r[vt[c]] for c in ben["khoa"]),
            "nhan_chi_tiet": tuple(r[vt[c]] for c in ben["chi_tiet"]),
            "gia_tri": _so(r[vt[ben["gia_tri"]]]),
        })
    return {"tep": p, "sheet": ten_sheet, "tieu_de": tieu_de, "dong": giu, "so_dong_bo": bo,
            "ly_do_bo": ly_do}


def _nhan(t) -> str:
    return " | ".join("" if x is None else str(x).strip() for x in t)


def _gop(spec, A, B):
    """Cộng theo khoá; trả về các dòng kết quả theo đúng thứ tự hiển thị."""
    def cong(ds):
        tong, chi, nhan = {}, {}, {}
        for d in ds:
            tong[d["khoa"]] = tong.get(d["khoa"], 0.0) + d["gia_tri"]
            nhan.setdefault(d["khoa"], d["nhan_khoa"])
            kc = (d["khoa"], d["chi_tiet"])
            chi[kc] = chi.get(kc, 0.0) + d["gia_tri"]
            nhan.setdefault(kc, d["nhan_chi_tiet"])
        return tong, chi, nhan

    ta, ca, na = cong(A["dong"])
    tb, cb, nb = cong(B["dong"])
    nhan = {**nb, **na}
    nguong = spec["ket_qua"]["nguong"]
    thu_tu = [(chuan(x),) for x in spec["ket_qua"]["thu_tu"]]
    con_lai = sorted((set(ta) | set(tb)) - set(thu_tu), key=lambda k: _nhan(nhan[k]))
    ngoai_mau = [_nhan(nhan[k]) for k in con_lai] if thu_tu else []
    dong = []
    for k in thu_tu + con_lai:
        cap = []
        if spec["a"]["chi_tiet"]:
            for kc in sorted({kc for kc in set(ca) | set(cb) if kc[0] == k},
                             key=lambda kc: _nhan(nhan[kc])):
                va, vb = ca.get(kc, 0.0), cb.get(kc, 0.0)
                if abs(va - vb) > nguong:
                    cap.append({"chi_tiet": _nhan(nhan[kc]), "a": va, "b": vb})
        hien = nhan.get(k)
        if hien is None:
            hien = next(x for x in spec["ket_qua"]["thu_tu"] if (chuan(x),) == k),
        va, vb = ta.get(k, 0.0), tb.get(k, 0.0)
        dong.append({"nhan": hien, "a": va, "b": vb, "lech": abs(va - vb) > nguong,
                     "cap_lech": cap})
    return dong, ngoai_mau


def _tung_dong(spec, A, B):
    nguong = spec["ket_qua"]["nguong"]
    kho = {}
    for d in B["dong"]:
        kho.setdefault((d["khoa"], d["chi_tiet"]), []).append({**d, "_dung": False})
    ra = []
    for d in A["dong"]:
        ung = [x for x in kho.get((d["khoa"], d["chi_tiet"]), []) if not x["_dung"]]
        khop = next((x for x in ung if abs(x["gia_tri"] - d["gia_tri"]) <= nguong), None)
        if khop is not None:
            khop["_dung"] = True
            ra.append({**_dong_ra(d, d["gia_tri"], khop["gia_tri"]), "ket_qua": "KHỚP"})
        elif ung:
            ung[0]["_dung"] = True
            ra.append({**_dong_ra(d, d["gia_tri"], ung[0]["gia_tri"]), "ket_qua": "LỆCH",
                       "so_dong_b": ung[0]["so_dong"]})
        else:
            ra.append({**_dong_ra(d, d["gia_tri"], None), "ket_qua": "CHỈ CÓ A"})
    for ds in kho.values():
        for x in ds:
            if not x["_dung"]:
                ra.append({**_dong_ra(x, None, x["gia_tri"]), "ket_qua": "CHỈ CÓ B",
                           "so_dong_a": None, "so_dong_b": x["so_dong"]})
    return ra


def _dong_ra(d, va, vb) -> dict:
    return {"nhan": d["nhan_khoa"] + d["nhan_chi_tiet"], "a": va, "b": vb,
            "so_dong_a": d["so_dong"] if va is not None else None, "so_dong_b": None}


def chay(spec: dict, tep: dict, tuy_chon: dict) -> dict:
    if set(tep) != {"a", "b"}:
        raise LoiQuyTrinh("Quy trình doi-chieu cần --a <file> và --b <file>")
    A, B = _doc_ben(spec["a"], tep["a"]), _doc_ben(spec["b"], tep["b"])
    spec = {**spec, "_ky": tuy_chon.get("ky")}
    k = spec["ket_qua"]
    ra_dir = tuy_chon.get("ra_thu_muc")
    ra = (Path(ra_dir) if ra_dir else tep["a"].parent) / ten_ra(k["ra"], tep["a"],
                                                                tuy_chon.get("ky"))
    canh_bao = []
    if k["che_do"] == "gop":
        dong, ngoai_mau = _gop(spec, A, B)
        _ghi_gop(spec, dong, ra)
        if ngoai_mau:
            canh_bao.append(f"Khoá chưa có trong thứ tự mẫu, đã nối thêm ở cuối: {ngoai_mau}")
        tom = {"so_khoa": len(dong), "so_khoa_lech": sum(d["lech"] for d in dong),
               "tong_a": round(sum(d["a"] for d in dong), 2),
               "tong_b": round(sum(d["b"] for d in dong), 2),
               "lech": [{"khoa": _nhan(d["nhan"]), "a": d["a"], "b": d["b"],
                         "cap_lech": d["cap_lech"][:GIOI_HAN_CAP],
                         "so_cap_lech": len(d["cap_lech"])} for d in dong if d["lech"]]}
    else:
        dong = _tung_dong(spec, A, B)
        _ghi_tung_dong(spec, dong, ra)
        dem = {}
        for d in dong:
            dem[d["ket_qua"]] = dem.get(d["ket_qua"], 0) + 1
        tom = {"theo_ket_qua": dem}
    return {
        "ket_qua": [{"a": str(tep["a"]), "b": str(tep["b"]), "ra": str(ra),
                     "so_dong_a": len(A["dong"]), "so_dong_b": len(B["dong"]),
                     "da_loai_a": A["ly_do_bo"], "da_loai_b": B["ly_do_bo"], **tom}],
        "canh_bao": canh_bao,
        "_quan_sat": {},
        "_dau_vao": [{"vai": t, "tep": str(X["tep"]), "bam": bam(X["tep"]), "sheet": X["sheet"],
                      "tieu_de": [str(x) if x is not None else "" for x in X["tieu_de"]]}
                     for t, X in (("a", A), ("b", B))],
    }


def _tieu_de_gia_tri(spec, t):
    """Tiêu đề cột giá trị của một bên; `tieu_de` trong quy trình được dùng {ky}."""
    b = spec[t]
    if b["tieu_de"]:
        return b["tieu_de"].replace("{ky}", spec.get("_ky") or "")
    return f"{b['gia_tri']} {b['nhan']}"


def _ghi_gop(spec, dong, ra: Path) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    from _xuat import luu

    a, b, k = spec["a"], spec["b"], spec["ket_qua"]
    so_khoa = len(a["khoa"])
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(ten_sheet_hop_le(k["sheet"]))
    ws.append(list(a["khoa"]) + [_tieu_de_gia_tri(spec, "a"), _tieu_de_gia_tri(spec, "b"),
                                 f"Chênh lệch ({a['nhan']} − {b['nhan']})", k["cot_ghi_chu"]])
    from openpyxl.utils import get_column_letter as L

    ca, cb, cc = L(so_khoa + 1), L(so_khoa + 2), L(so_khoa + 3)
    for i, d in enumerate(dong, start=2):
        ghi_chu = "; ".join(
            f"{x['chi_tiet']}: {a['nhan']} {x['a']:g}, {b['nhan']} {x['b']:g}"
            for x in d["cap_lech"][:GIOI_HAN_CAP])
        if len(d["cap_lech"]) > GIOI_HAN_CAP:
            ghi_chu += f"; … và {len(d['cap_lech']) - GIOI_HAN_CAP} cặp khác"
        ws.append(list(d["nhan"]) + [d["a"], d["b"], f"={ca}{i}-{cb}{i}", ghi_chu])
    n = len(dong) + 1
    ws.append(["TỔNG CỘNG"] + [""] * (so_khoa - 1) +
              [f"=SUM({ca}2:{ca}{n})", f"=SUM({cb}2:{cb}{n})", f"=SUM({cc}2:{cc}{n})", ""])
    loai = []
    for t in ("a", "b"):
        if spec[t]["loai_tru"]:
            loai.append(f"{spec[t]['nhan']}: không tính " + "; ".join(
                q["mo_ta"] or f"quy tắc {q['ma']}" for q in spec[t]["loai_tru"]))
    ws.cell(n + 3, 1, f"Nguồn: {a['nhan']} cộng {a['gia_tri']} theo {', '.join(a['khoa'])}; "
                      f"{b['nhan']} cộng {b['gia_tri']} theo {', '.join(b['khoa'])}."
                      + (" " + ". ".join(loai) + "." if loai else "")).font = Font(italic=True)
    mong = Side(style="thin")
    vien = Border(left=mong, right=mong, top=mong, bottom=mong)
    do = PatternFill("solid", fgColor="F8D7DA")
    so_cot = so_khoa + 4
    for hang in ws.iter_rows(min_row=1, max_row=n + 1, max_col=so_cot):
        r = hang[0].row
        lech = 1 < r <= n and dong[r - 2]["lech"]
        for c in hang:
            c.border = vien
            c.font = Font(bold=r in (1, n + 1))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if so_khoa < c.column <= so_khoa + 3:
                c.number_format = "#,##0.##;-#,##0.##;0"
            if lech:
                c.fill = do
    for i in range(1, so_khoa + 1):
        ws.column_dimensions[L(i)].width = 60 if so_khoa == 1 else 30
    for c in (ca, cb, cc):
        ws.column_dimensions[c].width = 16
    ws.column_dimensions[L(so_cot)].width = 80
    ws.freeze_panes = "A2"
    luu(wb, ra)


def _ghi_tung_dong(spec, dong, ra: Path) -> None:
    import openpyxl
    from openpyxl.styles import PatternFill

    from _xuat import dam_dong, luu, sheet_du_lieu

    a, b, k = spec["a"], spec["b"], spec["ket_qua"]
    tieu = (list(a["khoa"]) + list(a["chi_tiet"]) +
            [_tieu_de_gia_tri(spec, "a"), _tieu_de_gia_tri(spec, "b"),
             f"Chênh lệch ({a['nhan']} − {b['nhan']})", "Kết quả",
             f"Dòng {a['nhan']}", f"Dòng {b['nhan']}"])
    viet = [d for d in dong if not k["chi_ghi_lech"] or d["ket_qua"] != "KHỚP"]
    hang = []
    for d in viet:
        va, vb = d["a"], d["b"]
        hang.append(list(d["nhan"]) + [va, vb, (va or 0) - (vb or 0), d["ket_qua"],
                                       d["so_dong_a"], d["so_dong_b"]])
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = sheet_du_lieu(wb, ten_sheet_hop_le(k["sheet"]), tieu, hang)
    do = PatternFill("solid", fgColor="F8D7DA")
    ck = len(tieu) - 2
    for r in range(2, len(hang) + 2):
        if ws.cell(r, ck).value != "KHỚP":
            ws.cell(r, ck).fill = do
    t = wb.create_sheet("TÓM TẮT")
    t.append(["Kết quả", "Số dòng"])
    dem = {}
    for d in dong:
        dem[d["ket_qua"]] = dem.get(d["ket_qua"], 0) + 1
    for kq in ("KHỚP", "LỆCH", "CHỈ CÓ A", "CHỈ CÓ B"):
        t.append([kq.replace(" A", f" {a['nhan']}").replace(" B", f" {b['nhan']}"),
                  dem.get(kq, 0)])
    if k["chi_ghi_lech"]:
        t.append([])
        t.append(["Sheet đối chiếu chỉ ghi các dòng không khớp."])
    dam_dong(t, 1)
    t.column_dimensions["A"].width = 30
    luu(wb, ra)
