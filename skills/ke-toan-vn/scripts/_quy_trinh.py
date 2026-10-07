"""Lõi chung của các quy trình định kỳ: đọc công thức (recipe), danh mục, quy tắc, bảng.

Một *quy trình* là một file TOML trong `<bộ nhớ>/quy-trinh/<tên>.toml` mô tả trọn một việc
kế toán làm lặp lại hằng tháng (tách sổ theo khoa, lọc bản kết xuất, đối chiếu hai nguồn).
Skill giữ động cơ; quy trình giữ quy tắc riêng của đơn vị. Xem references/quy-trinh.md.

Module này không ghi file kết quả nào — việc đó thuộc từng động cơ (_tach, _loc, _doi_chieu_bang).
"""

from __future__ import annotations

import fnmatch
import hashlib
import json
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from _chung import tim_home

LOAI_HOP_LE = ("tach", "loc", "doi-chieu")
PHEP_SO = ("bang", "khac", "bat_dau", "chua", "am", "rong", "thuoc_tap", "ngoai_tap")


class LoiQuyTrinh(Exception):
    """Lỗi do quy trình hoặc dữ liệu đầu vào — báo cho người dùng, không phải lỗi lập trình."""


class CanHoi(Exception):
    """Dừng lại để hỏi kế toán (giá trị mới, bố cục đổi). Mang theo nội dung cần hỏi."""

    def __init__(self, thong_diep: str, chi_tiet: dict):
        super().__init__(thong_diep)
        self.chi_tiet = chi_tiet


# --- Chuẩn hoá -----------------------------------------------------------------------

def chuan(v) -> str:
    """Chuẩn hoá một giá trị ô để so khớp.

    NFC (macOS ghi NFD), gộp mọi khoảng trắng kể cả xuống dòng trong tiêu đề, bỏ hoa/thường
    ("Xét nghiệm " có dấu cách cuối, "vitamin"/"Vitamin"). Số nguyên lưu dạng float (8.0) về
    "8" để MADOITUONG đọc từ Excel khớp với 8 viết trong quy trình. KHÔNG bỏ số 0 đầu: mã
    khoa/mã tài sản "042" khác "42".
    """
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v).casefold()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, datetime):
        v = v.date() if v.time() == datetime.min.time() else v
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    s = unicodedata.normalize("NFC", str(v))
    return re.sub(r"\s+", " ", s).strip().casefold()


def la_so(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def ten_sheet_hop_le(ten: str) -> str:
    """Excel cấm []:*?/\\ trong tên sheet và giới hạn 31 ký tự."""
    s = re.sub(r"[\[\]:*?/\\]", "-", str(ten)).strip() or "Sheet"
    return s[:31]


def doc_toml(p: Path) -> dict:
    try:
        import tomllib
    except ImportError as e:  # Python < 3.11
        raise LoiQuyTrinh("Quy trình cần Python 3.11 trở lên (tomllib). Cài python3.11+ rồi "
                          "chạy lại bootstrap.sh để dựng lại venv.") from e
    try:
        with p.open("rb") as f:
            return tomllib.load(f)
    except tomllib.TOMLDecodeError as e:
        raise LoiQuyTrinh(f"{p.name}: sai cú pháp TOML — {e}") from e


# --- Danh mục -------------------------------------------------------------------------

def doc_danh_muc(home: Path) -> dict:
    """Đọc `<bộ nhớ>/danh-muc.toml`. Không có file thì danh mục rỗng."""
    p = home / "danh-muc.toml"
    return doc_toml(p) if p.exists() else {}


def tap_gia_tri(danh_muc: dict, ten: str) -> list:
    tap = danh_muc.get("tap", {})
    if ten not in tap:
        raise LoiQuyTrinh(f"Danh mục không có tập {ten!r}. Các tập đang có: {sorted(tap)}")
    return list(tap[ten])


def ten_goi(danh_muc: dict, bang: str, ma: str, nguon: str) -> list:
    """Tên của mục `ma` trong bảng danh mục `bang` theo cách nguồn `nguon` gọi nó.

    Cùng một khoa, mỗi file gọi một kiểu: "KHS" ở file lương, "SINH HÓA MIỄN DỊCH" ở file
    hao phí, "K. SINH HÓA" ở sổ tài sản. Danh mục giữ ánh xạ đó một lần cho mọi quy trình.
    """
    bang_dm = danh_muc.get(bang)
    if not isinstance(bang_dm, dict) or ma not in bang_dm:
        raise LoiQuyTrinh(f"Danh mục {bang!r} không có mục {ma!r}")
    muc = bang_dm[ma]
    if nguon not in muc:
        raise LoiQuyTrinh(
            f"Mục {bang}.{ma} chưa có tên gọi cho nguồn {nguon!r} — thêm `{nguon} = \"...\"` "
            f"vào danh-muc.toml")
    gt = muc[nguon]
    return list(gt) if isinstance(gt, list) else [gt]


# --- Quy tắc ---------------------------------------------------------------------------

def _chuan_dieu_kien(cot: str, dk: dict, danh_muc: dict, vi_tri: str) -> dict:
    if not isinstance(dk, dict):
        raise LoiQuyTrinh(f"{vi_tri}: điều kiện của cột {cot!r} phải là bảng, ví dụ "
                          f"{{ bang = [\"...\"] }}")
    la = set(dk) - set(PHEP_SO)
    if la:
        raise LoiQuyTrinh(f"{vi_tri}: phép so không hỗ trợ {sorted(la)} ở cột {cot!r}. "
                          f"Dùng: {', '.join(PHEP_SO)}")
    ra = dict(dk)
    if "thuoc_tap" in ra:
        ra.setdefault("bang", [])
        ra["bang"] = list(ra["bang"]) + tap_gia_tri(danh_muc, ra.pop("thuoc_tap"))
    if "ngoai_tap" in ra:
        ra.setdefault("khac", [])
        ra["khac"] = list(ra["khac"]) + tap_gia_tri(danh_muc, ra.pop("ngoai_tap"))
    for k in ("bang", "khac", "bat_dau", "chua"):
        if k in ra:
            if not isinstance(ra[k], list):
                raise LoiQuyTrinh(f"{vi_tri}: `{k}` ở cột {cot!r} phải là danh sách")
            ra[k] = frozenset(chuan(x) for x in ra[k]) if k in ("bang", "khac") \
                else tuple(chuan(x) for x in ra[k])
    return ra


def chuan_quy_tac(qt: dict, danh_muc: dict, vi_tri: str) -> dict:
    """Kiểm và chuẩn hoá một quy tắc {ma, mo_ta?, khi, tru?}."""
    la = set(qt) - {"ma", "mo_ta", "khi", "tru"}
    if la:
        raise LoiQuyTrinh(f"{vi_tri}: khoá lạ {sorted(la)} (chỉ có ma, mo_ta, khi, tru)")
    if "ma" not in qt or "khi" not in qt:
        raise LoiQuyTrinh(f"{vi_tri}: quy tắc cần `ma` và `khi`")
    tru = qt.get("tru", [])
    if isinstance(tru, dict):
        tru = [tru]
    return {
        "ma": str(qt["ma"]),
        "mo_ta": qt.get("mo_ta"),
        "khi": {c: _chuan_dieu_kien(c, d, danh_muc, vi_tri) for c, d in qt["khi"].items()},
        "tru": [{c: _chuan_dieu_kien(c, d, danh_muc, vi_tri) for c, d in t.items()}
                for t in tru],
    }


def thoa(dk: dict, v) -> bool:
    """Một ô có thoả TẤT CẢ phép so trong điều kiện của cột không."""
    s = chuan(v)
    if "bang" in dk and s not in dk["bang"]:
        return False
    if "khac" in dk and s in dk["khac"]:
        return False
    if "bat_dau" in dk and not s.startswith(dk["bat_dau"]):
        return False
    if "chua" in dk and not any(x in s for x in dk["chua"]):
        return False
    if "am" in dk and (la_so(v) and v < 0) != bool(dk["am"]):
        return False
    if "rong" in dk and (s == "") != bool(dk["rong"]):
        return False
    return True


def thoa_tat_ca(dieu_kien: dict, lay) -> bool:
    return all(thoa(dk, lay(cot)) for cot, dk in dieu_kien.items())


def khop_quy_tac(qt: dict, lay) -> bool:
    """`khi`: mọi cột phải thoả. `tru`: chỉ cần MỘT nhóm ngoại lệ thoả trọn là miễn."""
    if not thoa_tat_ca(qt["khi"], lay):
        return False
    return not any(thoa_tat_ca(t, lay) for t in qt["tru"])


def cot_cua_quy_tac(qt: dict) -> set:
    cot = set(qt["khi"])
    for t in qt["tru"]:
        cot |= set(t)
    return cot


# --- Quy trình -------------------------------------------------------------------------

def thu_muc_quy_trinh(home: Path) -> Path:
    return home / "quy-trinh"


def tim_quy_trinh(ten_hoac_duong_dan: str, home: Path) -> Path:
    p = Path(ten_hoac_duong_dan).expanduser()
    if p.suffix == ".toml" and p.exists():
        return p
    q = thu_muc_quy_trinh(home) / f"{ten_hoac_duong_dan}.toml"
    if q.exists():
        return q
    co = sorted(x.stem for x in thu_muc_quy_trinh(home).glob("*.toml"))
    raise LoiQuyTrinh(f"Không có quy trình {ten_hoac_duong_dan!r} trong "
                      f"{thu_muc_quy_trinh(home)}. Đang có: {co or 'chưa có quy trình nào'}")


KHOA_CHUNG = {"ten", "mo_ta", "loai", "nhat_ky", "cau_hoi_mo", "theo_doi"}


def doc_quy_trinh(duong_dan: Path) -> dict:
    qt = doc_toml(duong_dan)
    if qt.get("loai") not in LOAI_HOP_LE:
        raise LoiQuyTrinh(f"{duong_dan.name}: `loai` phải là một trong {LOAI_HOP_LE}")
    qt.setdefault("ten", duong_dan.stem)
    qt["_duong_dan"] = str(duong_dan)
    return qt


def kiem_khoa(bang: dict, cho_phep: set, vi_tri: str) -> None:
    """Bắt lỗi gõ sai tên khoá — TOML không báo, quy tắc gõ sai sẽ lặng lẽ không chạy."""
    la = {k for k in bang if not k.startswith("_")} - cho_phep
    if la:
        raise LoiQuyTrinh(f"{vi_tri}: khoá không nhận ra {sorted(la)}. Khoá hợp lệ: "
                          f"{sorted(cho_phep)}")


# --- Đọc bảng --------------------------------------------------------------------------

def chon_sheet(ten_cac_sheet: list, muon: str | None, ten_file: str) -> str:
    """Khớp tên sheet bỏ qua khoảng trắng thừa: sổ tài sản thật đặt tên sheet là " 2026 "."""
    if muon is None:
        return ten_cac_sheet[0]
    if muon in ten_cac_sheet:
        return muon
    for s in ten_cac_sheet:
        if chuan(s) == chuan(muon):
            return s
    raise LoiQuyTrinh(f"{ten_file}: không có sheet {muon!r}. Các sheet: {ten_cac_sheet}")


def chi_so_cot(tieu_de: list) -> dict:
    """Tên cột (đã chuẩn hoá) -> chỉ số. Tên trùng thì giữ cột đầu tiên."""
    ra = {}
    for i, h in enumerate(tieu_de):
        k = chuan(h)
        if k and k not in ra:
            ra[k] = i
    return ra


def tra_cot(chi_muc: dict, cot: str, ten_file: str) -> int:
    """Tìm cột theo tên tiêu đề, hoặc theo chữ cái nếu viết dạng "$Z".

    Ưu tiên tên: bản kết xuất kỳ sau chèn thêm một cột thì chữ cái trượt mà tên vẫn đúng.
    """
    from openpyxl.utils import column_index_from_string

    if cot.startswith("$"):
        return column_index_from_string(cot[1:].upper()) - 1
    k = chuan(cot)
    if k not in chi_muc:
        raise LoiQuyTrinh(f"{ten_file}: không có cột {cot!r} ở dòng tiêu đề")
    return chi_muc[k]


def doc_bang(duong_dan: Path, sheet: str | None, dong_tieu_de: int = 1):
    """Đọc giá trị (đã tính) của một sheet: (tên sheet, tiêu đề, [(số dòng, giá trị...)]).

    Bỏ dòng rỗng hoàn toàn. Cắt cột rỗng cuối tiêu đề (read_only trả tới max_column, có
    file khai báo 16.384 cột).
    """
    import openpyxl

    if not duong_dan.exists():
        raise LoiQuyTrinh(f"Không thấy file {duong_dan}")
    wb = openpyxl.load_workbook(duong_dan, read_only=True, data_only=True)
    try:
        ten = chon_sheet(wb.sheetnames, sheet, duong_dan.name)
        ws = wb[ten]
        tieu_de = None
        dong = []
        for so, r in enumerate(ws.iter_rows(values_only=True), start=1):
            if so == dong_tieu_de:
                tieu_de = list(r)
                while tieu_de and tieu_de[-1] is None:
                    tieu_de.pop()
                continue
            if tieu_de is None:
                continue
            n = len(tieu_de)
            gt = list(r[:n]) + [None] * (n - len(r))
            if all(v is None for v in gt):
                continue
            dong.append((so, gt))
    finally:
        wb.close()
    if tieu_de is None:
        raise LoiQuyTrinh(f"{duong_dan.name}/{ten}: không có dòng tiêu đề {dong_tieu_de}")
    return ten, tieu_de, dong


def bam(duong_dan: Path) -> str:
    h = hashlib.sha256()
    with duong_dan.open("rb") as f:
        for khoi in iter(lambda: f.read(1 << 20), b""):
            h.update(khoi)
    return h.hexdigest()[:16]


# --- Lịch sử chạy và giá trị đã biết ----------------------------------------------------

def thu_muc_lich_su(home: Path) -> Path:
    return thu_muc_quy_trinh(home) / "lich-su"


def lan_chay_truoc(home: Path, ten: str) -> dict | None:
    p = thu_muc_lich_su(home) / f"{ten}.jsonl"
    if not p.exists():
        return None
    cuoi = None
    for dong in p.read_text(encoding="utf-8").splitlines():
        if dong.strip():
            cuoi = json.loads(dong)
    return cuoi


def ghi_lich_su(home: Path, ten: str, ban_ghi: dict) -> None:
    d = thu_muc_lich_su(home)
    d.mkdir(parents=True, exist_ok=True)
    with (d / f"{ten}.jsonl").open("a", encoding="utf-8") as f:
        f.write(json.dumps(ban_ghi, ensure_ascii=False, default=str) + "\n")


def so_bo_cuc(truoc: list | None, nay: list) -> dict | None:
    """So tiêu đề kỳ này với lần chạy trước. None = không đổi hoặc chưa có lần trước."""
    if truoc is None:
        return None
    a = [str(x) if x is not None else "" for x in truoc]
    b = [str(x) if x is not None else "" for x in nay]
    if a == b:
        return None
    return {
        "them": [x for x in b if x and x not in a],
        "bot": [x for x in a if x and x not in b],
        "doi_vi_tri": a != b and set(a) == set(b),
        "so_cot_truoc": len(a),
        "so_cot_nay": len(b),
    }


def doc_gia_tri_da_biet(home: Path, ten: str) -> dict:
    p = thu_muc_lich_su(home) / f"{ten}.gia-tri.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


def ghi_gia_tri_da_biet(home: Path, ten: str, gia_tri: dict) -> None:
    d = thu_muc_lich_su(home)
    d.mkdir(parents=True, exist_ok=True)
    (d / f"{ten}.gia-tri.json").write_text(
        json.dumps({k: sorted(v, key=chuan) for k, v in gia_tri.items()},
                   ensure_ascii=False, indent=2),
        encoding="utf-8")


def kiem_gia_tri_moi(home: Path, ten: str, quan_sat: dict, chap_nhan: bool,
                     ghi: bool = True) -> dict:
    """Giá trị lần đầu xuất hiện ở các cột `theo_doi`.

    `quan_sat`: {cột: {giá trị đã chuẩn hoá: giá trị gốc}}. File lưu giá trị gốc để kế toán
    đọc được; so sánh luôn theo giá trị đã chuẩn hoá.

    Đây là hàng rào chống "quy tắc lặng lẽ không còn đúng": tháng sau phần mềm sinh thêm một
    MADOITUONG hay một TENLOAIVP mới thì không quy tắc nào nói phải giữ hay bỏ nó. Lần chạy
    đầu tiên chỉ ghi nhận; từ lần sau gặp giá trị lạ thì dừng để hỏi kế toán, trừ khi đã được
    chấp nhận (`--chap-nhan-gia-tri-moi`).
    """
    da_biet = doc_gia_tri_da_biet(home, ten)
    lan_dau = not da_biet
    moi = {}
    for cot, gt in quan_sat.items():
        if cot not in da_biet:
            continue
        cu = {chuan(x) for x in da_biet[cot]}
        la = sorted(gt[k] for k in gt if k not in cu)
        if la:
            moi[cot] = la
    if moi and not chap_nhan:
        raise CanHoi("Có giá trị chưa từng gặp ở các cột đang theo dõi — hỏi kế toán giữ hay "
                     "bỏ trước khi chạy tiếp", {"gia_tri_moi": moi})
    gop = {}
    for cot, gt in quan_sat.items():
        theo_khoa = {chuan(x): x for x in da_biet.get(cot, [])}
        for k, v in gt.items():
            theo_khoa.setdefault(k, v)
        gop[cot] = list(theo_khoa.values())
    if ghi:
        ghi_gia_tri_da_biet(home, ten, gop)
    return {"lan_dau_ghi_nhan": lan_dau, "gia_tri_moi_da_chap_nhan": moi}


# --- Kỳ báo cáo ---------------------------------------------------------------------------

_KY = re.compile(r"^\s*T?\s*(\d{1,2})\s*[./-]\s*(\d{4})\s*$", re.IGNORECASE)
_KY_ISO = re.compile(r"^\s*(\d{4})-(\d{1,2})\s*$")


def phan_tich_ky(ky: str) -> dict:
    """"T09.2026", "9/2026", "2026-09" -> {thang: "9", thang2: "09", nam: "2026"}."""
    m = _KY.match(ky)
    if m:
        thang, nam = int(m.group(1)), m.group(2)
    else:
        m = _KY_ISO.match(ky)
        if not m:
            raise LoiQuyTrinh(f"Không hiểu kỳ {ky!r} — viết dạng T09.2026, 09/2026 hoặc 2026-09")
        nam, thang = m.group(1), int(m.group(2))
    if not 1 <= thang <= 12:
        raise LoiQuyTrinh(f"Tháng {thang} trong kỳ {ky!r} không hợp lệ")
    return {"thang": str(thang), "thang2": f"{thang:02d}", "nam": nam}


def dien_ky(gia_tri, ky: str | None):
    """Thay {thang}, {thang2}, {nam} trong mọi chuỗi của quy trình.

    Tên cột của sổ đổi theo tháng ("Tháng 8/2026", "Lãi tháng 08 VCB"), nên quy trình viết
    "Tháng {thang}/{nam}" và lấy tháng từ --ky. {ky} và {ten_file} để nguyên — chúng thuộc
    tên file kết quả.
    """
    if isinstance(gia_tri, dict):
        return {k: dien_ky(v, ky) for k, v in gia_tri.items()}
    if isinstance(gia_tri, list):
        return [dien_ky(v, ky) for v in gia_tri]
    if not isinstance(gia_tri, str) or not re.search(r"\{(thang2?|nam)\}", gia_tri):
        return gia_tri
    if not ky:
        raise LoiQuyTrinh(f"Quy trình dùng {{thang}}/{{nam}} (ở {gia_tri!r}) nên cần --ky, "
                          f"ví dụ --ky T09.2026")
    p = phan_tich_ky(ky)
    for k in ("thang2", "thang", "nam"):
        gia_tri = gia_tri.replace("{" + k + "}", p[k])
    return gia_tri


# --- Tiện ích chung cho động cơ ------------------------------------------------------------

def ten_ra(mau: str, nguon: Path | None, ky: str | None) -> str:
    """Thay {ten_file} (tên file nguồn không đuôi) và {ky} trong mẫu tên file kết quả."""
    ten = mau.replace("{ten_file}", nguon.stem if nguon else "")
    if "{ky}" in ten:
        if not ky:
            raise LoiQuyTrinh(f"Tên file kết quả {mau!r} cần --ky (ví dụ T09.2026)")
        ten = ten.replace("{ky}", ky)
    return ten


def khop_ten_file(mau: list, ten_file: str) -> bool:
    return any(fnmatch.fnmatch(ten_file.casefold(), m.casefold()) for m in mau)


def home_mac_dinh() -> Path:
    return tim_home()
