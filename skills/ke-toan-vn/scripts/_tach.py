"""Động cơ `tach`: tách các dòng của một sheet ra từng sheet theo giá trị một cột khoá.

Việc lặp lại nhiều nhất của kế toán khoa/phòng: sổ lương, hao phí, sổ tài sản... mỗi tháng
tách theo khoa. Mỗi sheet kết quả giữ nguyên vùng đầu (mọi dòng tới dòng tiêu đề, gồm cả dòng
SUBTOTAL, ô gộp, độ rộng và cột ẩn) rồi chép các dòng của nhóm, kèm một dòng TỔNG CỘNG.
Dòng tổng dùng SUBTOTAL(9, ...) chứ không SUM: vùng đầu của sổ thường có sẵn SUBTOTAL cả cột,
và SUBTOTAL bỏ qua SUBTOTAL khác — dùng SUM thì dòng đầu cộng trùng thành gấp đôi.

Công thức trong dòng dữ liệu:
  - công thức chỉ tham chiếu CHÍNH dòng đó (=SUM(K4:X4), =J4) được dời theo dòng mới, để kế
    toán vẫn bấm vào ô thấy nó cộng từ đâu;
  - công thức tham chiếu dòng khác, sheet khác hay workbook ngoài được thay bằng giá trị đã
    tính: dòng đã đổi chỗ thì tham chiếu đó trỏ sai, còn workbook ngoài thì không mang theo.

Làm việc ngay trong workbook nguồn (rồi xoá các sheet không cần) thay vì workbook mới: chép
định dạng trong cùng workbook nhanh và giữ đúng phông mặc định của file gốc.
"""

from __future__ import annotations

import re
from copy import copy
from pathlib import Path

from _quy_trinh import (
    KHOA_CHUNG, LoiQuyTrinh, bam, chi_so_cot, chon_sheet, chuan, chuan_quy_tac, kiem_khoa,
    la_so, ten_goi, ten_ra, ten_sheet_hop_le, thoa_tat_ca, tra_cot,
)

KHOA_NGUON = {"sheet", "dong_tieu_de", "cot_khoa", "cot_tong", "cot_stt", "dong_du_lieu_khi"}
KHOA_KET_QUA = {"ra", "tom_tat", "nhan_tong", "cot_nhan_tong", "kem_sheet_nguon"}
KHOA_NHOM = {"ten", "gia_tri", "ma"}

_O = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")


def kiem(qt: dict, danh_muc: dict) -> dict:
    kiem_khoa(qt, KHOA_CHUNG | {"nguon", "ket_qua", "nhom", "ten_goi"}, qt["ten"])
    nguon = qt.get("nguon", {})
    kiem_khoa(nguon, KHOA_NGUON, "[nguon]")
    if "cot_khoa" not in nguon:
        raise LoiQuyTrinh("[nguon] cần `cot_khoa` — cột dùng để tách")
    kq = qt.get("ket_qua", {})
    kiem_khoa(kq, KHOA_KET_QUA, "[ket_qua]")
    tg = qt.get("ten_goi")
    if tg is not None:
        kiem_khoa(tg, {"bang", "nguon"}, "[ten_goi]")
    if not qt.get("nhom"):
        raise LoiQuyTrinh("Quy trình tach cần ít nhất một [[nhom]]")
    nhom, ten_da_dung = [], set()
    for i, n in enumerate(qt["nhom"], 1):
        vt = f"[[nhom]] thứ {i}"
        kiem_khoa(n, KHOA_NHOM, vt)
        gt = list(n.get("gia_tri", []))
        for ma in n.get("ma", []):
            if tg is None:
                raise LoiQuyTrinh(f"{vt}: dùng `ma` thì cần khai báo [ten_goi] (bảng danh "
                                  f"mục và tên nguồn)")
            gt += ten_goi(danh_muc, tg["bang"], ma, tg["nguon"])
        if not gt:
            raise LoiQuyTrinh(f"{vt}: cần `gia_tri` hoặc `ma`")
        ten = ten_sheet_hop_le(n.get("ten") or gt[0])
        if ten in ten_da_dung or ten == "TÓM TẮT":
            raise LoiQuyTrinh(f"{vt}: tên sheet {ten!r} bị trùng")
        ten_da_dung.add(ten)
        nhom.append({"ten": ten, "gia_tri": gt, "_khoa": {chuan(x) for x in gt}})
    trung = {}
    for n in nhom:
        for k in n["_khoa"]:
            trung.setdefault(k, []).append(n["ten"])
    hai_noi = {k: v for k, v in trung.items() if len(v) > 1}
    if hai_noi:
        raise LoiQuyTrinh(f"Một giá trị nằm trong hai nhóm: {hai_noi}")
    return {
        **qt,
        "nguon": {"dong_tieu_de": 1, "cot_tong": [], **nguon},
        "ket_qua": {"ra": "{ten_file}-TACH.xlsx", "tom_tat": True, "nhan_tong": "TỔNG CỘNG",
                    "kem_sheet_nguon": False, **kq},
        "nhom": nhom,
        "_dong_du_lieu_khi": chuan_quy_tac(
            {"ma": "_", "khi": nguon["dong_du_lieu_khi"]}, danh_muc, "[nguon]")["khi"]
        if nguon.get("dong_du_lieu_khi") else None,
    }


def _cong_thuc_trong_dong(f: str, dong: int, cot_cam: frozenset = frozenset()) -> bool:
    """Công thức chỉ tham chiếu ô của chính dòng `dong`, cùng sheet, không cố định dòng.

    `cot_cam`: cột mà giá trị ở sheet mới khác nguồn (cột STT được đánh lại). Công thức đọc
    cột đó phải giữ giá trị cũ — sổ tài sản ghép STT vào mã tài sản ("57-21133 - MMTBCD..."),
    đánh số lại mà để công thức sống là đổi luôn mã.
    """
    if "!" in f or "[" in f:
        return False
    from openpyxl.formula import Tokenizer

    try:
        tok = Tokenizer(f)
    except Exception:  # noqa: BLE001 — công thức lạ thì coi như không dời được, lấy giá trị
        return False
    co_o = False
    for t in tok.items:
        if t.type != "OPERAND" or t.subtype != "RANGE":
            continue
        if re.fullmatch(r"\$?[A-Z]{1,3}:\$?[A-Z]{1,3}", t.value):
            continue
        o = _O.findall(t.value)
        if not o or _O.sub("", t.value).replace(":", ""):
            return False
        for _, chu, co_dinh_dong, so in o:
            if co_dinh_dong or int(so) != dong:
                return False
        if cot_cam:
            from openpyxl.utils import column_index_from_string as idx

            dau, _, cuoi = t.value.replace("$", "").partition(":")
            c1 = idx(_O.match(dau).group(2))
            c2 = idx(_O.match(cuoi).group(2)) if cuoi else c1
            if any(c1 <= c <= c2 for c in cot_cam):
                return False
        co_o = True
    return co_o


def _gia_tri_o(o_ct, o_gt, dong_cu: int, dong_moi: int, cot_cam: frozenset):
    from openpyxl.formula.translate import Translator

    v = o_ct.value
    if isinstance(v, str) and v.startswith("=") and _cong_thuc_trong_dong(v, dong_cu, cot_cam):
        return Translator(v, origin=f"A{dong_cu}").translate_formula(f"A{dong_moi}")
    return o_gt.value


def chay(spec: dict, tep: list[Path], tuy_chon: dict) -> dict:
    if len(tep) != 1:
        raise LoiQuyTrinh("Quy trình tach nhận đúng một file nguồn")
    import openpyxl
    from openpyxl.utils import get_column_letter

    p = tep[0]
    n, k = spec["nguon"], spec["ket_qua"]
    if not p.exists():
        raise LoiQuyTrinh(f"Không thấy file {p}")
    wb = openpyxl.load_workbook(p)
    wb_gt = openpyxl.load_workbook(p, data_only=True)
    ten_sheet = chon_sheet(wb.sheetnames, n.get("sheet"), p.name)
    ws, ws_gt = wb[ten_sheet], wb_gt[ten_sheet]
    hd = n["dong_tieu_de"]

    tieu_de = [ws_gt.cell(hd, c).value for c in range(1, ws.max_column + 1)]
    while tieu_de and tieu_de[-1] is None:
        tieu_de.pop()
    # Cột cuối có dữ liệu thật, ở bất kỳ dòng nào: cột ghi chú không có tiêu đề (hao phí có
    # cột K "CẤN TRỪ T7...") vẫn phải mang theo. Không dùng max_column — sổ tài sản khai báo
    # 16.384 cột chỉ vì định dạng. Đọc thẳng các ô đã tồn tại thay vì quét cả lưới.
    so_cot = max([len(tieu_de)] + [c for (_, c), o in ws_gt._cells.items() if o.value is not None])
    cm = chi_so_cot(tieu_de)
    i_khoa = tra_cot(cm, n["cot_khoa"], p.name)
    i_tong = [tra_cot(cm, c, p.name) for c in n["cot_tong"]]
    i_stt = tra_cot(cm, n["cot_stt"], p.name) if n.get("cot_stt") else None
    dl_khi = spec["_dong_du_lieu_khi"]
    vi_tri_dk = {c: tra_cot(cm, c, p.name) for c in (dl_khi or {})}

    theo_khoa = {}
    for nh in spec["nhom"]:
        for x in nh["_khoa"]:
            theo_khoa[x] = nh["ten"]
    dong_nhom = {nh["ten"]: [] for nh in spec["nhom"]}
    ngoai_nhom, khong_phai_dl, gap = {}, 0, set()
    quan_sat = {c: {} for c in spec.get("theo_doi", [])}
    vi_tri_td = {c: tra_cot(cm, c, p.name) for c in quan_sat}
    for r in range(hd + 1, ws.max_row + 1):
        hang = [ws_gt.cell(r, c + 1).value for c in range(so_cot)]
        if all(v is None for v in hang):
            continue
        if dl_khi and not thoa_tat_ca(dl_khi, lambda c, h=hang: h[vi_tri_dk[c]]):
            khong_phai_dl += 1
            continue
        for c, i in vi_tri_td.items():
            quan_sat[c].setdefault(chuan(hang[i]), str(hang[i] or "").strip())
        kh = chuan(hang[i_khoa])
        if kh in theo_khoa:
            dong_nhom[theo_khoa[kh]].append(r)
            gap.add(kh)
        else:
            nhan = str(hang[i_khoa] or "").strip() or "(trống)"
            ngoai_nhom[nhan] = ngoai_nhom.get(nhan, 0) + 1

    nguon_ws = ws
    tong_nguon = {}
    for nh in spec["nhom"]:
        tong_nguon[nh["ten"]] = [
            round(sum(v for r in dong_nhom[nh["ten"]]
                      if la_so(v := ws_gt.cell(r, i + 1).value)), 2) for i in i_tong]

    cot_cam = frozenset({i_stt + 1}) if i_stt is not None else frozenset()
    merges_dau = [str(m) for m in nguon_ws.merged_cells.ranges if m.max_row <= hd]
    tong_dong = {}
    for nh in spec["nhom"]:
        ten = nh["ten"]
        moi = wb.create_sheet(ten)
        for chu, dim in nguon_ws.column_dimensions.items():
            moi.column_dimensions[chu].width = dim.width
            moi.column_dimensions[chu].hidden = dim.hidden
        for r in range(1, hd + 1):
            for c in range(1, so_cot + 1):
                o = nguon_ws.cell(r, c)
                d = moi.cell(r, c)
                v = o.value
                if isinstance(v, str) and v.startswith("=") and ("!" in v or "[" in v):
                    v = ws_gt.cell(r, c).value
                d.value = v
                if o.has_style:
                    d._style = copy(o._style)
            if nguon_ws.row_dimensions[r].height:
                moi.row_dimensions[r].height = nguon_ws.row_dimensions[r].height
        for m in merges_dau:
            moi.merge_cells(m)
        for j, r in enumerate(dong_nhom[ten]):
            rd = hd + 1 + j
            for c in range(1, so_cot + 1):
                o = nguon_ws.cell(r, c)
                d = moi.cell(rd, c)
                d.value = _gia_tri_o(o, ws_gt.cell(r, c), r, rd, cot_cam)
                if o.has_style:
                    d._style = copy(o._style)
            if i_stt is not None:
                moi.cell(rd, i_stt + 1).value = j + 1
            if nguon_ws.row_dimensions[r].height:
                moi.row_dimensions[rd].height = nguon_ws.row_dimensions[r].height
        cuoi = hd + len(dong_nhom[ten])
        if i_tong:
            from openpyxl.styles import Font

            rt = cuoi + 1
            if k.get("cot_nhan_tong"):
                c_nhan = tra_cot(cm, k["cot_nhan_tong"], p.name)
            else:
                c_nhan = min(i_tong) - 1 if min(i_tong) > 0 and min(i_tong) - 1 not in i_tong \
                    else 0
            o = moi.cell(rt, c_nhan + 1, k["nhan_tong"])
            o.font = Font(bold=True)
            for i in i_tong:
                chu = get_column_letter(i + 1)
                o = moi.cell(rt, i + 1)
                o.value = f"=SUBTOTAL(9,{chu}{hd + 1}:{chu}{cuoi})" if cuoi > hd else 0
                o.font = Font(bold=True)
                o.number_format = nguon_ws.cell(hd + 1, i + 1).number_format or "#,##0"
            tong_dong[ten] = rt
        moi.freeze_panes = nguon_ws.freeze_panes
        moi.auto_filter.ref = f"A{hd}:{get_column_letter(so_cot)}{max(cuoi, hd)}"

    if k["tom_tat"]:
        _tom_tat(wb, spec, dong_nhom, tong_dong, tong_nguon, n["cot_tong"], i_tong)

    giu = {nh["ten"] for nh in spec["nhom"]} | {"TÓM TẮT"}
    if k["kem_sheet_nguon"]:
        giu.add(ten_sheet)
    for s in list(wb.sheetnames):
        if s not in giu:
            del wb[s]
    _don_dep_workbook(wb)
    if k["kem_sheet_nguon"]:
        wb.move_sheet(ten_sheet, offset=-wb.sheetnames.index(ten_sheet))

    ra_dir = tuy_chon.get("ra_thu_muc")
    ra = (Path(ra_dir) if ra_dir else p.parent) / ten_ra(k["ra"], p, tuy_chon.get("ky"))
    from _xuat import luu

    luu(wb, ra)

    khong_gap = sorted(x for nh in spec["nhom"] for x in nh["gia_tri"] if chuan(x) not in gap)
    return {
        "ket_qua": [{
            "nguon": str(p), "ra": str(ra), "sheet": ten_sheet,
            "nhom": [{"sheet": nh["ten"], "gia_tri": nh["gia_tri"],
                      "so_dong": len(dong_nhom[nh["ten"]]),
                      "tong": dict(zip(n["cot_tong"], tong_nguon[nh["ten"]]))}
                     for nh in spec["nhom"]],
            "dong_khong_phai_du_lieu": khong_phai_dl,
            "gia_tri_khong_co_trong_du_lieu": khong_gap,
            "dong_ngoai_cac_nhom": dict(sorted(ngoai_nhom.items(), key=lambda kv: -kv[1])[:40]),
        }],
        "canh_bao": ([f"Giá trị trong quy trình không có dòng nào kỳ này: {khong_gap} — kiểm "
                      f"tra khoa/phòng có đổi tên không"] if khong_gap else []),
        "_quan_sat": quan_sat,
        "_dau_vao": [{"tep": str(p), "bam": bam(p), "sheet": ten_sheet,
                      "tieu_de": [str(x) if x is not None else "" for x in tieu_de]}],
    }


def _tom_tat(wb, spec, dong_nhom, tong_dong, tong_nguon, ten_cot_tong, i_tong) -> None:
    """Sheet TÓM TẮT: tổng từng nhóm lấy bằng công thức từ chính sheet nhóm, đặt cạnh tổng
    tính thẳng từ file nguồn — lệch là biết ngay bước chép có vấn đề."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet("TÓM TẮT")
    ws.append([f"Quy trình: {spec['ten']}"])
    ws.append([])
    tieu = ["Sheet", "Giá trị gộp", "Số dòng"]
    for c in ten_cot_tong:
        tieu += [f"{c} (sheet)", f"{c} (nguồn)"]
    ws.append(tieu)
    for c in ws[3]:
        c.font = Font(bold=True)
    for nh in spec["nhom"]:
        ten = nh["ten"]
        hang = [ten, ", ".join(nh["gia_tri"]), len(dong_nhom[ten])]
        for j, i in enumerate(i_tong):
            chu = get_column_letter(i + 1)
            hang += [f"='{ten}'!{chu}{tong_dong[ten]}", tong_nguon[ten][j]]
        ws.append(hang)
    dau, cuoi = 4, 3 + len(spec["nhom"])
    tong = ["TỔNG CỘNG", "", f"=SUM(C{dau}:C{cuoi})"]
    for j in range(len(i_tong) * 2):
        chu = get_column_letter(4 + j)
        tong.append(f"=SUM({chu}{dau}:{chu}{cuoi})")
    ws.append(tong)
    for c in ws[cuoi + 1]:
        c.font = Font(bold=True)
    if i_tong:
        lech = "+".join(f"ABS({get_column_letter(4 + 2 * j)}{cuoi + 1}-"
                        f"{get_column_letter(5 + 2 * j)}{cuoi + 1})" for j in range(len(i_tong)))
        ws.append([])
        ws.append(["Kết luận", f'=IF({lech}<0.5,"Khớp","LỆCH")'])
        ws.cell(cuoi + 3, 1).font = Font(bold=True)
    for j in range(len(i_tong) * 2):
        for r in range(dau, cuoi + 2):
            ws.cell(r, 4 + j).number_format = "#,##0"
        ws.column_dimensions[get_column_letter(4 + j)].width = 20
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 36


def _don_dep_workbook(wb) -> None:
    """Gỡ tên vùng trỏ vào sheet đã xoá hoặc workbook ngoài, để Excel không đòi sửa file."""
    con = set(wb.sheetnames)
    for ten in list(wb.defined_names):
        dn = wb.defined_names[ten]
        try:
            hong = "[" in (dn.attr_text or "") or any(s not in con for s, _ in dn.destinations)
        except Exception:  # noqa: BLE001 — tên vùng dạng công thức lạ: bỏ cho an toàn
            hong = True
        if hong:
            del wb.defined_names[ten]
    for ws in wb.worksheets:
        for ten in list(ws.defined_names):
            del ws.defined_names[ten]
    wb._external_links = []
