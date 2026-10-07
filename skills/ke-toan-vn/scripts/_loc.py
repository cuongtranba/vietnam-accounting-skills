"""Động cơ `loc`: lọc từng dòng của một bảng thành phần GIỮ và phần LOẠI BỎ có ghi lý do.

Mỗi file nguồn sinh một workbook gồm:
  <sheet_giu>  các dòng giữ, nguyên thứ tự cột (để bước sau trỏ cột theo chữ cái vẫn đúng)
  <sheet_bo>   các dòng bị loại, cột cuối ghi quy tắc nào loại dòng đó
  TÓM TẮT      đối chiếu bằng công thức: gốc = giữ + loại, cả số dòng lẫn tổng cột tiền

Một dòng bị loại nếu dính BẤT KỲ quy tắc nào. Biến thể (`bien_the`) cho phép một nhóm file
— nhận theo tên file — có thêm ngoại lệ hoặc bỏ hẳn một quy tắc.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from _quy_trinh import (
    KHOA_CHUNG, LoiQuyTrinh, bam, chi_so_cot, chuan, chuan_quy_tac, cot_cua_quy_tac,
    doc_bang, khop_quy_tac, khop_ten_file, kiem_khoa, la_so, ten_ra, ten_sheet_hop_le,
    thoa_tat_ca, tra_cot,
)

KHOA_NGUON = {"sheet", "dong_tieu_de", "cot_bat_buoc", "cot_tien", "dong_du_lieu_khi"}
KHOA_KET_QUA = {"ra", "sheet_giu", "sheet_bo", "cot_ly_do"}
KHOA_BIEN_THE = {"ten", "mo_ta", "ap_khi_ten_file", "them_tru", "bo_quy_tac"}


def kiem(qt: dict, danh_muc: dict) -> dict:
    kiem_khoa(qt, KHOA_CHUNG | {"nguon", "ket_qua", "quy_tac", "bien_the"}, qt["ten"])
    nguon = qt.get("nguon", {})
    kiem_khoa(nguon, KHOA_NGUON, "[nguon]")
    kq = qt.get("ket_qua", {})
    kiem_khoa(kq, KHOA_KET_QUA, "[ket_qua]")
    if not qt.get("quy_tac"):
        raise LoiQuyTrinh("Quy trình loc cần ít nhất một [[quy_tac]]")
    quy_tac = [chuan_quy_tac(q, danh_muc, f"[[quy_tac]] thứ {i}")
               for i, q in enumerate(qt["quy_tac"], 1)]
    ma = [q["ma"] for q in quy_tac]
    if len(set(ma)) != len(ma):
        raise LoiQuyTrinh(f"Mã quy tắc bị trùng: {ma}")
    bien_the = []
    for i, b in enumerate(qt.get("bien_the", []), 1):
        vt = f"[[bien_the]] thứ {i}"
        kiem_khoa(b, KHOA_BIEN_THE, vt)
        if "ten" not in b:
            raise LoiQuyTrinh(f"{vt}: cần `ten`")
        them = {}
        for m, ds in b.get("them_tru", {}).items():
            if m not in ma:
                raise LoiQuyTrinh(f"{vt}: them_tru nhắc quy tắc {m!r} không tồn tại")
            ds = [ds] if isinstance(ds, dict) else ds
            them[m] = [chuan_quy_tac({"ma": m, "khi": t}, danh_muc, vt)["khi"] for t in ds]
        for m in b.get("bo_quy_tac", []):
            if str(m) not in ma:
                raise LoiQuyTrinh(f"{vt}: bo_quy_tac nhắc quy tắc {m!r} không tồn tại")
        bien_the.append({**b, "them_tru": them,
                         "bo_quy_tac": [str(m) for m in b.get("bo_quy_tac", [])],
                         "ap_khi_ten_file": list(b.get("ap_khi_ten_file", []))})
    return {
        **qt,
        "nguon": {"dong_tieu_de": 1, **nguon},
        "ket_qua": {"ra": "LOC-{ten_file}.xlsx", "sheet_giu": "GIỮ LẠI",
                    "sheet_bo": "LOẠI BỎ", "cot_ly_do": "LÝ DO LOẠI", **kq},
        "quy_tac": quy_tac,
        "bien_the": bien_the,
        "_dong_du_lieu_khi": chuan_quy_tac(
            {"ma": "_", "khi": nguon["dong_du_lieu_khi"]}, danh_muc, "[nguon]")["khi"]
        if nguon.get("dong_du_lieu_khi") else None,
    }


def ap_bien_the(spec: dict, ten_file: str, ep: str | None, tat: bool) -> tuple[list, str | None]:
    """Bộ quy tắc áp cho một file, sau khi cộng biến thể (nếu có)."""
    chon = None
    if not tat:
        for b in spec["bien_the"]:
            if (ep and chuan(b["ten"]) == chuan(ep)) or \
               (not ep and khop_ten_file(b["ap_khi_ten_file"], ten_file)):
                chon = b
                break
        if ep and chon is None:
            raise LoiQuyTrinh(f"Không có biến thể {ep!r}")
    if chon is None:
        return spec["quy_tac"], None
    ra = []
    for q in spec["quy_tac"]:
        if q["ma"] in chon["bo_quy_tac"]:
            continue
        ra.append({**q, "tru": q["tru"] + chon["them_tru"].get(q["ma"], [])})
    return ra, chon["ten"]


def ly_do(q: dict, lay) -> str:
    if q["mo_ta"]:
        return f"Quy tắc {q['ma']}: {q['mo_ta']}"
    chi_tiet = ", ".join(f"{c} {str(lay(c)).strip()}" for c in q["khi"])
    return f"Quy tắc {q['ma']}: {chi_tiet}"


def doc_va_phan_loai(spec: dict, duong_dan: Path, ep_bien_the, tat_bien_the) -> dict:
    n = spec["nguon"]
    ten_sheet, tieu_de, dong = doc_bang(duong_dan, n.get("sheet"), n["dong_tieu_de"])
    cm = chi_so_cot(tieu_de)
    quy_tac, bien_the = ap_bien_the(spec, duong_dan.name, ep_bien_the, tat_bien_the)
    can = set(n.get("cot_bat_buoc", [])) | set(spec.get("theo_doi", []))
    for q in spec["quy_tac"] + quy_tac:
        can |= cot_cua_quy_tac(q)
    if n.get("cot_tien"):
        can.add(n["cot_tien"])
    thieu = [c for c in sorted(can) if not c.startswith("$") and chuan(c) not in cm]
    if thieu:
        raise LoiQuyTrinh(f"{duong_dan.name}/{ten_sheet}: thiếu cột {thieu}. Bố cục bản kết "
                          f"xuất có thể đã đổi — xem lại dòng tiêu đề {n['dong_tieu_de']}.")
    vi_tri = {c: tra_cot(cm, c, duong_dan.name) for c in can}
    dl_khi = spec["_dong_du_lieu_khi"]

    giu, bo, khong_phai_dl, dem = [], [], 0, {}
    quan_sat = {c: {} for c in spec.get("theo_doi", [])}
    for _, r in dong:
        def lay(c, r=r):
            return r[vi_tri[c]]
        if dl_khi and not thoa_tat_ca(dl_khi, lay):
            khong_phai_dl += 1
            continue
        for c in quan_sat:
            v = lay(c)
            quan_sat[c].setdefault(chuan(v), "" if v is None else str(v).strip())
        ld = [ly_do(q, lay) for q in quy_tac if khop_quy_tac(q, lay)]
        if not ld:
            giu.append(r)
            continue
        bo.append(r + ["; ".join(ld)])
        for x in ld:
            ma = x.split(":", 1)[0]
            dem[ma] = dem.get(ma, 0) + 1
    return {"duong_dan": duong_dan, "sheet": ten_sheet, "tieu_de": tieu_de, "cm": cm,
            "giu": giu, "bo": bo, "dem": dem, "bien_the": bien_the,
            "khong_phai_du_lieu": khong_phai_dl, "quan_sat": quan_sat,
            "quy_tac_ap": [q["ma"] for q in quy_tac]}


def ghi(spec: dict, kq: dict, ra: Path) -> None:
    import openpyxl
    from openpyxl.utils import get_column_letter

    from _xuat import dam_dong, luu, sheet_du_lieu

    k = spec["ket_qua"]
    tieu_de = kq["tieu_de"]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sg, sb = ten_sheet_hop_le(k["sheet_giu"]), ten_sheet_hop_le(k["sheet_bo"])
    sheet_du_lieu(wb, sg, tieu_de, kq["giu"])
    ws_bo = sheet_du_lieu(wb, sb, tieu_de + [k["cot_ly_do"]], kq["bo"])
    ld = get_column_letter(len(tieu_de) + 1)
    ws_bo.column_dimensions[ld].width = 60

    ws = wb.create_sheet("TÓM TẮT")
    g, b = f"'{sg}'", f"'{sb}'"
    cot_tien = spec["nguon"].get("cot_tien")
    dong = [
        ["Quy trình", spec["ten"]],
        ["Nguồn", f"{kq['duong_dan'].name} / sheet {kq['sheet']}"],
        ["Biến thể", kq["bien_the"] or "Không"],
        ["Ngày lập", datetime.now().strftime("%d/%m/%Y %H:%M")],
        [],
    ]
    if cot_tien:
        i = kq["cm"][chuan(cot_tien)]
        st = get_column_letter(i + 1)
        tong = sum(r[i] for r in kq["giu"] + kq["bo"] if la_so(r[i]))
        dong += [
            ["", "Số dòng", f"Tổng {cot_tien}"],
            ["File gốc", len(kq["giu"]) + len(kq["bo"]), tong],
            [f"Giữ lại ({sg})", f"=COUNTA({g}!A:A)-1", f"=SUM({g}!{st}:{st})"],
            [f"Loại bỏ ({sb})", f"=COUNTA({b}!A:A)-1", f"=SUM({b}!{st}:{st})"],
            ["Gốc − giữ − loại", "=B7-B8-B9", "=C7-C8-C9"],
            ["Kết luận", '=IF(AND(B10=0,ABS(C10)<0.5),"Khớp","LỆCH")'],
        ]
    else:
        dong += [
            ["", "Số dòng"],
            ["File gốc", len(kq["giu"]) + len(kq["bo"])],
            [f"Giữ lại ({sg})", f"=COUNTA({g}!A:A)-1"],
            [f"Loại bỏ ({sb})", f"=COUNTA({b}!A:A)-1"],
            ["Gốc − giữ − loại", "=B7-B8-B9"],
            ["Kết luận", '=IF(B10=0,"Khớp","LỆCH")'],
        ]
    dong += [[], ["Số dòng dính từng quy tắc (một dòng có thể dính nhiều quy tắc)"]]
    for r in dong:
        ws.append(r)
    for q in spec["quy_tac"]:
        nhan = f"Quy tắc {q['ma']}" + (f" — {q['mo_ta']}" if q["mo_ta"] else "")
        ws.append([nhan, f'=COUNTIF({b}!{ld}:{ld},"*Quy tắc {q["ma"]}:*")'])
    dam_dong(ws, 6, 13)
    for r in range(7, 11):
        ws.cell(r, 3).number_format = "#,##0"
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 20
    luu(wb, ra)


def chay(spec: dict, tep: list[Path], tuy_chon: dict) -> dict:
    """Đọc và phân loại MỌI file trước, rồi mới ghi: gặp giá trị lạ thì không file nào bị ghi."""
    if not tep:
        raise LoiQuyTrinh("Cần ít nhất một file nguồn")
    ket = [doc_va_phan_loai(spec, p, tuy_chon.get("bien_the"), tuy_chon.get("khong_bien_the"))
           for p in tep]
    quan_sat = {}
    for k in ket:
        for c, v in k["quan_sat"].items():
            quan_sat.setdefault(c, {}).update(v)
    return {"_ket": ket, "_quan_sat": quan_sat}


def ghi_tat_ca(spec: dict, trung_gian: dict, tuy_chon: dict) -> dict:
    ra_dir = tuy_chon.get("ra_thu_muc")
    ket_qua, dau_vao = [], []
    for k in trung_gian["_ket"]:
        p = k["duong_dan"]
        ra = (Path(ra_dir) if ra_dir else p.parent) / ten_ra(spec["ket_qua"]["ra"], p,
                                                             tuy_chon.get("ky"))
        ghi(spec, k, ra)
        i = k["cm"].get(chuan(spec["nguon"].get("cot_tien") or "")) \
            if spec["nguon"].get("cot_tien") else None

        def tong(ds, i=i):
            return round(sum(r[i] for r in ds if la_so(r[i])), 2) if i is not None else None
        ket_qua.append({
            "nguon": str(p), "ra": str(ra), "sheet": k["sheet"], "bien_the": k["bien_the"],
            "so_dong_goc": len(k["giu"]) + len(k["bo"]),
            "so_dong_giu": len(k["giu"]), "so_dong_bo": len(k["bo"]),
            "dong_khong_phai_du_lieu": k["khong_phai_du_lieu"],
            "tong_tien_giu": tong(k["giu"]), "tong_tien_bo": tong([r[:-1] for r in k["bo"]]),
            "ly_do_bo": dict(sorted(k["dem"].items(), key=lambda kv: -kv[1])),
        })
        dau_vao.append({"tep": str(p), "bam": bam(p), "sheet": k["sheet"],
                        "tieu_de": [str(x) if x is not None else "" for x in k["tieu_de"]]})
    return {"ket_qua": ket_qua, "_dau_vao": dau_vao}
