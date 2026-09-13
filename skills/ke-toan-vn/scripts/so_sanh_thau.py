#!/usr/bin/env python3
"""Lập bảng so sánh giá đánh giá giữa các nhà thầu (gói mua sắm hàng hóa).

    so_sanh_thau.py hsdt.json --danh-muc danh_muc_hsmt.csv --ra so_sanh.xlsx
    so_sanh_thau.py hsdt.json --hsmt hsmt.pdf --ra so_sanh.xlsx

Đầu vào là kết quả của doc_hsdt.py, cộng danh mục hàng hóa yêu cầu trong HSMT.

CÔNG THỨC (NĐ 214/2025/NĐ-CP):

    G = (giá dự thầu ± sửa lỗi ± hiệu chỉnh sai lệch) − giảm giá + ΔG + ΔƯĐ

Ba cái bẫy số học được xử lý đúng ở đây:

  1. Thư giảm giá trừ SAU CÙNG. Sửa lỗi và hiệu chỉnh sai lệch tính trên giá dự thầu CHƯA trừ
     giảm giá; tỷ lệ % sai lệch thiếu cũng so với giá ghi trong đơn dự thầu. Trừ sớm là ra sai.
  2. Chào thiếu thuế/phí KHÔNG tính vào sai lệch thiếu.
  3. Sai lệch thiếu không quá 10% giá dự thầu là điều kiện xét duyệt trúng thầu — vượt thì
     cảnh báo, nhưng KHÔNG tự tuyên bố loại.

HAI CHUẨN ĐƠN GIÁ NGƯỢC CHIỀU CHO CÙNG MỘT HẠNG MỤC CHÀO THIẾU (Điều 31 NĐ 214/2025):
  - Điểm c khoản 2 — hiệu chỉnh để SO SÁNH, XẾP HẠNG (cột "Giá đánh giá (G)"): đơn giá CAO NHẤT
    trong các HSDT khác đã vượt bước kỹ thuật → dự toán → giá gói thầu.
  - Khoản 4 — áp đơn giá cho GIÁ ĐỀ NGHỊ TRÚNG THẦU (giá hợp đồng) của riêng nhà thầu đang xếp
    hạng nhất, nếu HSDT của họ không có đơn giá phần chào thiếu: đơn giá THẤP NHẤT trong các HSDT
    khác đã vượt bước kỹ thuật → dự toán → giá gói thầu.
  "Giá đề nghị trúng thầu" theo Luật VẪN bao gồm giá trị hiệu chỉnh sai lệch (định nghĩa "giá đề
  nghị trúng thầu" = giá dự thầu đã sửa lỗi, hiệu chỉnh sai lệch, trừ giảm giá) — chỉ đơn giá dùng
  để tính lại khác đi. Cái thật sự không cộng vào giá hợp đồng là ΔG/ΔƯĐ ở trên.

NGUYÊN TẮC: đây là bảng tính SƠ BỘ. Script không kết luận nhà thầu nào đạt hay trúng thầu.
Xếp hạng in ra là xếp hạng theo con số, không phải kết quả lựa chọn nhà thầu.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _chung import can_thu_vien, canh_bao, in_json  # noqa: E402
from chuan_hoa import lam_sach_text, so_vn  # noqa: E402

NGUONG_SAI_LECH = Decimal("0.10")   # 10% giá dự thầu

# Chào thiếu các khoản này KHÔNG tính vào sai lệch thiếu
_THUE_PHI = re.compile(r"\b(thuế|phí|lệ\s*phí|vat|gtgt)\b", re.IGNORECASE)

LUU_Y = ("KẾT QUẢ TÍNH TOÁN SƠ BỘ. Tổ chuyên gia phải rà soát, đối chiếu với HSMT và HSDT gốc "
         "trước khi kết luận. Tài liệu này KHÔNG phải là kết quả đánh giá hồ sơ dự thầu.")


def chuan_ten(s: str) -> str:
    """Chuẩn hoá tên hàng để so khớp giữa HSMT và HSDT."""
    s = unicodedata.normalize("NFD", (s or "").lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").replace("đ", "d")
    s = re.sub(r"[^a-z0-9]+", " ", s).strip()
    return s


def khop_ten(ten: str, ung_vien: list[str]) -> str | None:
    """Khớp tên hàng: trùng khớp trước, rồi tới chứa nhau, rồi tới trùng nhiều từ."""
    c = chuan_ten(ten)
    if not c:
        return None
    kho = {chuan_ten(u): u for u in ung_vien}
    if c in kho:
        return kho[c]
    for k, goc in kho.items():
        if k and (k in c or c in k):
            return goc
    tu_c = set(c.split())
    tot, diem_tot = None, 0
    for k, goc in kho.items():
        chung = len(tu_c & set(k.split()))
        if chung > diem_tot and chung >= 2:
            tot, diem_tot = goc, chung
    return tot


def doc_danh_muc_hsmt(duong_dan: Path) -> list[dict]:
    """Đọc danh mục hàng hóa yêu cầu từ CSV/Excel hoặc PDF HSMT."""
    import pandas as pd

    if duong_dan.suffix.lower() == ".pdf":
        import pdfplumber
        muc = []
        with pdfplumber.open(duong_dan) as pdf:
            for tr in pdf.pages:
                for bang in (tr.extract_tables() or []):
                    if len(bang) < 2:
                        continue
                    tieu_de = " ".join(str(o or "") for o in bang[0]).lower()
                    if not re.search(r"(tên|hàng\s*hóa|mô\s*tả)", tieu_de):
                        continue
                    for hang in bang[1:]:
                        o = [lam_sach_text(x, "") or "" for x in hang]
                        ten = next((x for x in o if x and len(x) > 2
                                    and not re.fullmatch(r"[\d.,\s]*", x)), None)
                        if ten and not re.match(r"^(tổng|cộng)", ten, re.IGNORECASE):
                            sl = next((so_vn(x) for x in o if so_vn(x) is not None), None)
                            muc.append({"ten": ten, "so_luong": sl})
        return muc

    khung = (pd.read_csv(duong_dan, dtype=str) if duong_dan.suffix.lower() == ".csv"
             else pd.read_excel(duong_dan, dtype=str))
    khung.columns = [lam_sach_text(c, "") for c in khung.columns]

    # Khớp tên cột KHÔNG phân biệt dấu: tiêu đề cột trong file thật hay là 'ten_hang_hoa',
    # 'TenHang', 'so_luong' — viết không dấu. Đòi đúng dấu thì không khớp gì cả và rơi về
    # cột đầu tiên (thường là STT), làm sai toàn bộ việc đối chiếu danh mục.
    def tim_cot(*tu_khoa):
        for c in khung.columns:
            can = chuan_ten(c)
            if any(tk in can for tk in tu_khoa):
                return c
        return None

    cot_ten = tim_cot("ten hang", "ten hh", "hang hoa", "mo ta", "ten") or khung.columns[0]
    cot_sl = tim_cot("so luong", "sl")
    return [{"ten": r[cot_ten], "so_luong": so_vn(r[cot_sl]) if cot_sl else None}
            for _, r in khung.iterrows() if lam_sach_text(r[cot_ten], "")]


def _ghi_luu_y(tep: Path, luu_y: str, so_cot: int) -> None:
    """Đặt câu lưu ý lên dòng 1 sheet Tổng hợp, in đậm và đỏ, để không ai bỏ qua."""
    import openpyxl
    from openpyxl.styles import Alignment, Font

    wb = openpyxl.load_workbook(tep)
    ws = wb["Tong hop"]
    ws["A1"] = luu_y
    ws["A1"].font = Font(bold=True, color="C8102E")
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(so_cot, 2))
    ws.row_dimensions[1].height = 46
    ws.freeze_panes = "A4"
    for c in range(1, so_cot + 1):
        ws.cell(3, c).font = Font(bold=True)
    wb.save(tep)


def doc_ket_qua_ky_thuat(duong_dan, ten_nha_thau: list[str]):
    """Đọc phiếu chấm kỹ thuật đã điền → (vượt, trượt, chưa rõ).

    Quy tắc: chỉ xếp một nhà thầu vào nhóm VƯỢT khi MỌI tiêu chí của họ đã được kết luận Đạt.
    Chỉ cần một ô "Không đạt" là trượt; còn ô trống nghĩa là chưa chấm xong — không suy diễn
    thành đạt, vì suy diễn ở đây sẽ âm thầm làm đổi giá trị hiệu chỉnh và đổi thứ hạng.
    """
    if not duong_dan:
        return set(), set(), set()

    import pandas as pd
    p = Path(duong_dan).expanduser()
    if p.suffix.lower() == ".csv":
        khung = pd.read_csv(p, dtype=str)
    else:
        # cham_ky_thuat.py ghi dòng lưu ý ở đầu nên header thật nằm ở dòng thứ 3
        khung = pd.read_excel(p, sheet_name="Phieu cham ky thuat", dtype=str, header=2)
    khung.columns = [chuan_ten(c) for c in khung.columns]

    c_nt = next((c for c in khung.columns if "nha thau" in c), None)
    c_kl = next((c for c in khung.columns if "ket luan" in c), None)
    if not c_nt or not c_kl:
        raise SystemExit(f"Phiếu chấm kỹ thuật phải có cột 'nha_thau' và 'ket_luan'. "
                         f"Đang có: {list(khung.columns)}")

    trang_thai: dict[str, set[str]] = {t: set() for t in ten_nha_thau}
    for _, r in khung.iterrows():
        nt = lam_sach_text(r[c_nt], "")
        if not nt:
            continue
        khop = next((t for t in ten_nha_thau if chuan_ten(t) == chuan_ten(nt)), nt)
        kl = chuan_ten(r[c_kl]) if r[c_kl] is not None else ""
        trang_thai.setdefault(khop, set()).add(
            "dat" if kl == "dat" else "khong dat" if kl.startswith("khong") else "trong")

    vuot, truot, chua_ro = set(), set(), set()
    for nt, cac in trang_thai.items():
        if not cac:
            chua_ro.add(nt)
        elif "khong dat" in cac:
            truot.add(nt)
        elif "trong" in cac:
            chua_ro.add(nt)
        else:
            vuot.add(nt)
    return vuot, truot, chua_ro


def _gt(truong):
    """Lấy giá trị từ cấu trúc {gia_tri, nguon_trang, ...} của doc_hsdt.py."""
    if isinstance(truong, dict):
        return truong.get("gia_tri")
    return truong


def _trang(truong):
    return truong.get("nguon_trang") if isinstance(truong, dict) else None


def don_gia_dong_hang(d: dict) -> Decimal | None:
    """Đơn giá của một dòng hàng — đọc trực tiếp, hoặc suy từ thành tiền/số lượng nếu thiếu.

    Điểm d khoản 1 Điều 31 NĐ 214/2025: nếu cột thành tiền đã điền nhưng thiếu đơn giá tương
    ứng, đơn giá = thành tiền / số lượng. Đây là bước SỬA LỖI, phải làm trước khi coi dòng này
    là "không có đơn giá" — nếu không, đơn giá cao/thấp nhất tính hụt mất một mốc lẽ ra dùng
    được, và một cạnh tranh giá thật bị loại khỏi phép so sánh mà không ai biết.

    Trả None khi không suy được (thiếu cả đơn giá lẫn thành tiền, hoặc thiếu/bằng 0 số lượng) —
    gọi nơi dùng phải coi đây là "chưa xác định", không tự ý giả định bằng 0 hay 1.
    """
    dg = d.get("don_gia")
    if dg is not None:
        return Decimal(str(dg))
    tt, sl = d.get("thanh_tien"), d.get("so_luong")
    if tt is None or sl is None:
        return None
    # so_luong có thể tới dưới dạng chuỗi "0" sau khi qua JSON (doc_hsdt.py tuần tự hoá Decimal
    # thành str) — chuỗi "0" vẫn truthy trong Python, nên phải parse ra Decimal rồi so sánh thật
    # sự bằng 0, không chỉ kiểm truthiness, kẻo chia cho 0 làm sập cả báo cáo.
    sl_so = Decimal(str(sl))
    if sl_so == 0:
        return None
    return Decimal(str(tt)) / sl_so


def main() -> int:
    p = argparse.ArgumentParser(
        description="Lập bảng so sánh giá đánh giá giữa các nhà thầu.",
        epilog="Kết quả là bảng tính sơ bộ, KHÔNG phải kết quả đánh giá HSDT.")
    p.add_argument("hsdt_json", help="File JSON do doc_hsdt.py sinh ra")
    p.add_argument("--danh-muc", help="Danh mục hàng hóa yêu cầu (CSV/Excel)")
    p.add_argument("--hsmt", help="File HSMT (PDF) — thử đọc danh mục từ bảng trong đó")
    p.add_argument("--du-toan", type=float,
                   help="Đơn giá dự toán, dùng khi HSDT thiếu đơn giá cho phần sai lệch thiếu")
    p.add_argument("--ket-qua-ky-thuat",
                   help="Phiếu chấm kỹ thuật đã được tổ chuyên gia điền cột ket_luan "
                        "(từ cham_ky_thuat.py). Quyết định nhà thầu nào thuộc nhóm 'vượt bước "
                        "đánh giá về kỹ thuật' — nhóm này mới được lấy đơn giá để hiệu chỉnh.")
    p.add_argument("--ra", help="File Excel kết quả")
    args = p.parse_args()

    can_thu_vien("pandas", "openpyxl")
    import pandas as pd

    du_lieu = json.loads(Path(args.hsdt_json).expanduser().read_text(encoding="utf-8"))
    nha_thau = du_lieu.get("nha_thau", du_lieu if isinstance(du_lieu, list) else [])
    if not nha_thau:
        print("Không có dữ liệu nhà thầu trong file JSON", file=sys.stderr)
        return 1

    yeu_cau = []
    if args.danh_muc:
        yeu_cau = doc_danh_muc_hsmt(Path(args.danh_muc).expanduser())
    elif args.hsmt:
        yeu_cau = doc_danh_muc_hsmt(Path(args.hsmt).expanduser())

    canh_bao_chung = []
    if not yeu_cau:
        canh_bao_chung.append(
            "Không có danh mục hàng hóa yêu cầu từ HSMT — KHÔNG kiểm được chào thiếu/chào thừa. "
            "Cung cấp --danh-muc để kiểm đầy đủ.")

    ten_yeu_cau = [m["ten"] for m in yeu_cau]

    # --- Nhóm vượt bước đánh giá về kỹ thuật ---------------------------------
    ten_nha_thau = [nt["nha_thau"] for nt in nha_thau]
    vuot_kt, truot_kt, chua_ro = doc_ket_qua_ky_thuat(args.ket_qua_ky_thuat, ten_nha_thau)

    if args.ket_qua_ky_thuat is None:
        vuot_kt = set(ten_nha_thau)
        canh_bao_chung.append(
            "CHƯA có kết quả đánh giá kỹ thuật. Điểm c khoản 2 Điều 31 NĐ 214/2025 yêu cầu lấy "
            "đơn giá cao nhất trong các HSDT ĐÃ VƯỢT bước kỹ thuật; đang tạm lấy trên TẤT CẢ "
            "nhà thầu, nên giá trị hiệu chỉnh có thể sai. Chạy cham_ky_thuat.py, để tổ chuyên gia "
            "điền cột ket_luan, rồi chạy lại với --ket-qua-ky-thuat.")
    else:
        if chua_ro:
            canh_bao_chung.append(
                f"Chưa điền kết luận kỹ thuật cho: {', '.join(sorted(chua_ro))}. "
                "Các nhà thầu này KHÔNG được tính vào nhóm lấy đơn giá hiệu chỉnh — "
                "điền nốt cột ket_luan rồi chạy lại để có số đúng.")
        if truot_kt:
            canh_bao_chung.append(
                f"Không vượt bước kỹ thuật (theo kết luận của tổ chuyên gia): "
                f"{', '.join(sorted(truot_kt))}. Đơn giá của các nhà thầu này KHÔNG dùng để "
                "hiệu chỉnh sai lệch, và họ không vào danh sách xếp hạng.")
        if not vuot_kt:
            canh_bao_chung.append(
                "KHÔNG có nhà thầu nào vượt bước kỹ thuật — không xếp hạng được.")

    # Đơn giá CAO NHẤT của từng mặt hàng trong các HSDT khác — dùng để hiệu chỉnh phần chào thiếu.
    #
    # Điểm c khoản 2 Điều 31 NĐ 214/2025: "lấy mức đơn giá chào CAO NHẤT đối với hạng mục này
    # trong số các hồ sơ dự thầu khác vượt qua bước đánh giá về kỹ thuật để làm cơ sở hiệu chỉnh
    # sai lệch". Rất dễ nhầm thành 'thấp nhất' vì nhiều tài liệu đang lưu hành vẫn chép theo
    # NĐ 63/2014 đã hết hiệu lực. Dùng nhầm chiều là đảo thứ hạng nhà thầu.
    #
    # Chỉ lấy trên các HSDT ĐÃ VƯỢT bước kỹ thuật — điều luật nói rõ như vậy, và nó có ý nghĩa
    # thực chất: một nhà thầu trượt kỹ thuật thì đơn giá của họ không phải mặt bằng hợp lệ.
    don_gia_cao_nhat: dict[str, Decimal] = {}
    for nt in nha_thau:
        if nt["nha_thau"] not in vuot_kt:
            continue
        for d in nt.get("danh_muc", []):
            khop = khop_ten(d["ten"], ten_yeu_cau) if ten_yeu_cau else d["ten"]
            if not khop:
                continue
            dg = don_gia_dong_hang(d)
            if dg is not None and (khop not in don_gia_cao_nhat or dg > don_gia_cao_nhat[khop]):
                don_gia_cao_nhat[khop] = dg

    # Đơn giá THẤP NHẤT của từng mặt hàng — dùng RIÊNG cho khoản 4 Điều 31 NĐ 214/2025: áp cho
    # nhà thầu ĐANG XẾP HẠNG NHẤT khi tính giá đề nghị trúng thầu (giá đưa vào hợp đồng), nếu
    # HSDT của họ không có đơn giá cho phần chào thiếu. NGƯỢC CHIỀU với don_gia_cao_nhat ở trên —
    # hai biến phục vụ hai mục đích khác nhau (xếp hạng vs. giá hợp đồng thật), đừng gộp lại.
    don_gia_thap_nhat: dict[str, Decimal] = {}
    for nt in nha_thau:
        if nt["nha_thau"] not in vuot_kt:
            continue
        for d in nt.get("danh_muc", []):
            khop = khop_ten(d["ten"], ten_yeu_cau) if ten_yeu_cau else d["ten"]
            if not khop:
                continue
            dg = don_gia_dong_hang(d)
            if dg is not None and (khop not in don_gia_thap_nhat or dg < don_gia_thap_nhat[khop]):
                don_gia_thap_nhat[khop] = dg

    bang_tong, bang_doi_chieu, bang_sua_loi, bang_canh_bao, bang_can_cu = [], [], [], [], []
    thieu_theo_nha_thau: dict[str, list[dict]] = {}
    khong_xac_dinh_theo_nha_thau: dict[str, list[str]] = {}

    for nt in nha_thau:
        ten_nt = nt["nha_thau"]
        gia_du_thau = _gt(nt.get("gia_du_thau"))
        gia_du_thau = Decimal(str(gia_du_thau)) if gia_du_thau is not None else None
        giam_gia = _gt(nt.get("gia_giam"))
        giam_gia = Decimal(str(giam_gia)) if giam_gia is not None else Decimal(0)

        for truong, nhan in (("gia_du_thau", "Giá dự thầu"), ("gia_giam", "Thư giảm giá"),
                             ("bao_lanh_du_thau", "Bảo đảm dự thầu"),
                             ("bao_lanh_hieu_luc_den", "Hiệu lực bảo đảm"),
                             ("doanh_thu_binh_quan", "Doanh thu bình quân"),
                             ("mst", "Mã số thuế")):
            o = nt.get(truong)
            bang_can_cu.append({
                "Nhà thầu": ten_nt, "Chỉ tiêu": nhan,
                "Giá trị": _gt(o), "Nguồn": _trang(o),
                "Trích dẫn": (o.get("trich_dan") if isinstance(o, dict) else None),
                "Ghi chú": (o.get("ghi_chu") if isinstance(o, dict) else None),
            })

        # --- Sửa lỗi số học --------------------------------------------------
        tong_sua_loi = Decimal(0)
        for d in nt.get("danh_muc", []):
            loi = d.get("loi_so_hoc")
            if loi:
                chenh = Decimal(str(loi["chenh_lech"]))
                tong_sua_loi -= chenh   # đưa về giá trị đúng
                bang_sua_loi.append({
                    "Nhà thầu": ten_nt, "Mặt hàng": d["ten"],
                    "Số lượng": d.get("so_luong"), "Đơn giá": d.get("don_gia"),
                    "Thành tiền ghi trong HSDT": loi["ghi_trong_hsdt"],
                    "Thành tiền tính đúng": loi["tinh_ra"],
                    "Giá trị sửa lỗi": -chenh,
                    "Nguồn": d.get("nguon_trang"),
                })

        # --- Hiệu chỉnh sai lệch --------------------------------------------
        da_chao = {}
        for d in nt.get("danh_muc", []):
            khop = khop_ten(d["ten"], ten_yeu_cau) if ten_yeu_cau else None
            if khop:
                da_chao[khop] = d

        sai_lech_thieu = Decimal(0)
        sai_lech_thua = Decimal(0)

        for muc in yeu_cau:
            ten = muc["ten"]
            if ten in da_chao:
                d_item = da_chao[ten]
                dg_goc = d_item.get("don_gia")
                dg_rieng = don_gia_dong_hang(d_item)
                bang_doi_chieu.append({
                    "Mặt hàng (HSMT)": ten, "Nhà thầu": ten_nt, "Tình trạng": "Có chào",
                    "Đơn giá": dg_rieng,
                    "Thành tiền": d_item.get("thanh_tien"),
                    "Nguồn": d_item.get("nguon_trang"),
                })
                # Mặt hàng có tên trong bảng giá dự thầu, nên KHÔNG vào sai_lech_thieu/
                # thieu_theo_nha_thau (nó không phải "không được liệt kê"). Nhưng nếu bảng giá
                # thiếu đơn giá riêng, phải suy từ thành tiền/số lượng trước (điểm d khoản 1 Điều
                # 31 — don_gia_dong_hang() đã làm việc đó); chỉ khi vẫn không suy được mới coi là
                # thật sự không có giá, và phải nói rõ ra thay vì để một cột trống không lý do.
                if dg_goc is None and dg_rieng is not None:
                    bang_canh_bao.append({
                        "Nhà thầu": ten_nt, "Loại": "Đơn giá suy từ thành tiền",
                        "Nội dung": (
                            f"'{ten}' không có đơn giá ghi trực tiếp trong HSDT — đã suy ra "
                            f"{dg_rieng} bằng thành tiền ({d_item.get('thanh_tien')}) chia số "
                            f"lượng, theo điểm d khoản 1 Điều 31 NĐ 214/2025. Kiểm lại trang "
                            f"{d_item.get('nguon_trang') or '(không rõ)'} trước khi dùng."
                        ),
                    })
                elif dg_rieng is None:
                    tt_rieng = d_item.get("thanh_tien")
                    ly_do = ("KHÔNG đọc được đơn giá lẫn thành tiền" if tt_rieng is None else
                             f"có thành tiền ({tt_rieng}) nhưng thiếu số lượng nên không suy được "
                             "đơn giá")
                    bang_canh_bao.append({
                        "Nhà thầu": ten_nt, "Loại": "Thiếu đơn giá dòng đã chào",
                        "Nội dung": (
                            f"'{ten}' có trong bảng giá dự thầu nhưng {ly_do}, nên KHÔNG được "
                            "tính vào sai lệch thiếu (khoản 2/4 Điều 31) ở bảng này. Kiểm tra lại "
                            "trang HSDT trước khi dùng số liệu."
                        ),
                    })
                continue

            # Chào thiếu. Thuế/phí thì KHÔNG tính vào sai lệch thiếu.
            if _THUE_PHI.search(ten):
                bang_doi_chieu.append({
                    "Mặt hàng (HSMT)": ten, "Nhà thầu": ten_nt,
                    "Tình trạng": "Chào thiếu (thuế/phí — KHÔNG tính sai lệch)",
                    "Đơn giá": None, "Thành tiền": None, "Nguồn": None,
                })
                continue

            dg = don_gia_cao_nhat.get(ten)
            nguon_dg = ("đơn giá CAO NHẤT trong các HSDT khác vượt bước kỹ thuật "
                        "(điểm c khoản 2 Điều 31 NĐ 214/2025 — ưu tiên 1)")
            if dg is None and args.du_toan:
                dg, nguon_dg = Decimal(str(args.du_toan)), "đơn giá dự toán gói thầu (ưu tiên 2)"

            # KHÔNG mặc định số lượng = 1 khi thiếu — đó là bịa số cho một phép nhân, y hệt lỗi
            # "đừng điền 0" mà skill này luôn tránh. Thiếu số lượng thì để None và cảnh báo.
            sl = Decimal(str(muc["so_luong"])) if muc.get("so_luong") else None
            gia_tri = dg * sl if (dg is not None and sl is not None) else None
            if gia_tri is not None:
                sai_lech_thieu += gia_tri
            thieu_theo_nha_thau.setdefault(ten_nt, []).append(
                {"ten": ten, "so_luong": muc.get("so_luong")})

            bang_doi_chieu.append({
                "Mặt hàng (HSMT)": ten, "Nhà thầu": ten_nt, "Tình trạng": "CHÀO THIẾU",
                "Đơn giá": dg, "Thành tiền": gia_tri,
                "Nguồn": nguon_dg if gia_tri is not None else "KHÔNG có đơn giá và/hoặc số lượng — cần xác định",
            })
            if gia_tri is None:
                khong_xac_dinh_theo_nha_thau.setdefault(ten_nt, []).append(ten)
                thieu = ([] if dg is not None else ["đơn giá"]) + ([] if sl is not None else ["số lượng yêu cầu (HSMT)"])
                bang_canh_bao.append({
                    "Nhà thầu": ten_nt, "Loại": "Thiếu dữ liệu để hiệu chỉnh",
                    "Nội dung": (
                        f"Chào thiếu '{ten}' nhưng KHÔNG xác định được {' và '.join(thieu)} nên "
                        "chưa tính được giá trị hiệu chỉnh (khoản 2 Điều 31 NĐ 214/2025). "
                        + ("Đơn giá theo thứ tự ưu tiên: CAO NHẤT trong các HSDT khác vượt bước "
                           "kỹ thuật → đơn giá dự toán gói thầu → đơn giá hình thành giá gói thầu — "
                           "truyền --du-toan hoặc xác định thủ công. " if dg is None else "")
                        + ("Kiểm tra lại cột số lượng trong danh mục HSMT." if sl is None else "")
                    ),
                })

        for d in nt.get("danh_muc", []):
            if ten_yeu_cau and not khop_ten(d["ten"], ten_yeu_cau):
                tt = d.get("thanh_tien")
                if tt is not None:
                    sai_lech_thua += Decimal(str(tt))
                bang_doi_chieu.append({
                    "Mặt hàng (HSMT)": "(không có trong HSMT)", "Nhà thầu": ten_nt,
                    "Tình trạng": "CHÀO THỪA", "Đơn giá": d.get("don_gia"),
                    "Thành tiền": tt, "Nguồn": d.get("nguon_trang"),
                })

        hieu_chinh = sai_lech_thieu - sai_lech_thua

        # --- Giá đánh giá ----------------------------------------------------
        # Thứ tự đúng: sửa lỗi và hiệu chỉnh trên giá CHƯA trừ giảm giá, rồi mới trừ giảm giá.
        if gia_du_thau is not None:
            sau_sua_loi = gia_du_thau + tong_sua_loi
            sau_hieu_chinh = sau_sua_loi + hieu_chinh
            gia_danh_gia = sau_hieu_chinh - giam_gia
            # % sai lệch thiếu so với giá dự thầu ghi trong ĐƠN dự thầu (chưa trừ giảm giá)
            ty_le = (sai_lech_thieu / gia_du_thau) if gia_du_thau else None
        else:
            sau_sua_loi = sau_hieu_chinh = gia_danh_gia = ty_le = None
            bang_canh_bao.append({
                "Nhà thầu": ten_nt, "Loại": "Thiếu giá dự thầu",
                "Nội dung": "Không trích được giá dự thầu — không tính được giá đánh giá.",
            })

        if ty_le is not None and ty_le > NGUONG_SAI_LECH:
            bang_canh_bao.append({
                "Nhà thầu": ten_nt, "Loại": "Sai lệch thiếu vượt ngưỡng",
                "Nội dung": (f"Sai lệch thiếu {sai_lech_thieu:,.0f} = {ty_le * 100:.2f}% giá dự "
                             f"thầu, vượt ngưỡng 10% — là một điều kiện xét duyệt trúng thầu. "
                             "Tổ chuyên gia xem xét, script không kết luận loại."),
            })

        hl = _gt(nt.get("bao_lanh_hieu_luc_den"))
        if hl:
            bang_canh_bao.append({
                "Nhà thầu": ten_nt, "Loại": "Bảo đảm dự thầu",
                "Nội dung": (f"Bảo đảm dự thầu có hiệu lực đến {hl} "
                             f"(nguồn: {_trang(nt.get('bao_lanh_hieu_luc_den'))}). "
                             "Đối chiếu với thời hạn yêu cầu trong HSMT."),
            })

        for c in nt.get("canh_bao", []):
            bang_canh_bao.append({"Nhà thầu": ten_nt, "Loại": "Trích xuất", "Nội dung": c})

        tt_kt = ("vượt kỹ thuật" if ten_nt in vuot_kt else
                 "KHÔNG vượt kỹ thuật" if ten_nt in truot_kt else
                 "chưa có kết luận kỹ thuật")

        bang_tong.append({
            "Nhà thầu": ten_nt,
            "Kỹ thuật": tt_kt,
            "Giá dự thầu": gia_du_thau,
            "Giá trị sửa lỗi": tong_sua_loi,
            "Sai lệch thiếu": sai_lech_thieu,
            "Sai lệch thừa": sai_lech_thua,
            "Giá trị hiệu chỉnh sai lệch": hieu_chinh,
            "Sau sửa lỗi + hiệu chỉnh": sau_hieu_chinh,
            "Thư giảm giá": giam_gia,
            "% sai lệch thiếu": (float(ty_le * 100) if ty_le is not None else None),
            "Giá đánh giá (G)": gia_danh_gia,
            "Giá đề nghị trúng thầu (dự kiến, khoản 4 Điều 31)": None,
            "Ghi chú": ("ΔG và ΔƯĐ chưa tính — chỉ tính khi HSMT có quy định công thức"),
        })

    # Chỉ xếp hạng nhà thầu đã vượt bước kỹ thuật — bước 5 của quy trình đứng SAU bước 3,
    # nên nhà thầu trượt kỹ thuật không có mặt trong danh sách xếp hạng.
    #
    # VÀ: không xếp hạng nhà thầu còn hạng mục chào thiếu chưa định giá được (thiếu đơn giá
    # và/hoặc số lượng, xem canh_bao "Thiếu dữ liệu để hiệu chỉnh"). "Giá đánh giá (G)" của họ
    # hiện đang ngầm coi phần chưa định giá đó bằng 0 — con số chỉ mang tính tham khảo, KHÔNG đủ
    # tin cậy để xếp hạng cạnh các nhà thầu khác, vì giá trị thật của phần thiếu có thể đổi thứ
    # hạng. Bổ sung dữ liệu rồi chạy lại mới có xếp hạng.
    co_gia = [r for r in bang_tong
              if r["Giá đánh giá (G)"] is not None and r["Nhà thầu"] in vuot_kt
              and r["Nhà thầu"] not in khong_xac_dinh_theo_nha_thau]
    co_gia.sort(key=lambda r: r["Giá đánh giá (G)"])
    for i, r in enumerate(co_gia, 1):
        r["Xếp hạng sơ bộ"] = i
    for r in bang_tong:
        r.setdefault("Xếp hạng sơ bộ", None)

    for ten_nt, muc in khong_xac_dinh_theo_nha_thau.items():
        if ten_nt in vuot_kt:
            bang_canh_bao.append({
                "Nhà thầu": ten_nt, "Loại": "Chưa xếp hạng — sai lệch chưa định giá được",
                "Nội dung": (
                    f"{ten_nt} CHƯA được xếp hạng vì còn {len(muc)} hạng mục chào thiếu chưa "
                    f"định giá được ({', '.join(muc)}) — 'Giá đánh giá (G)' hiện tính coi phần đó "
                    "bằng 0, có thể sai và đổi thứ hạng thật khi bổ sung dữ liệu. Xác định đơn giá/"
                    "số lượng còn thiếu (xem cảnh báo 'Thiếu dữ liệu để hiệu chỉnh' ở trên) rồi "
                    "chạy lại."
                ),
            })

    # --- Khoản 4 Điều 31: giá đề nghị trúng thầu dự kiến, riêng cho nhà thầu xếp hạng nhất -----
    # Chỉ áp dụng cho nhà thầu ĐANG xếp hạng nhất (dù xếp theo giá thấp nhất hay giá đánh giá G),
    # và chỉ cho phần chào thiếu mà HSDT của họ hoàn toàn không có đơn giá riêng. Đơn giá dùng ở
    # đây là don_gia_thap_nhat — NGƯỢC CHIỀU với đơn giá đã dùng để xếp hạng nhà thầu này.
    gia_de_nghi_khoan4 = None
    if co_gia:
        nguoi_thang = co_gia[0]["Nhà thầu"]
        muc_thieu = thieu_theo_nha_thau.get(nguoi_thang, [])
        if muc_thieu:
            gia_tri_khoan4 = Decimal(0)
            thieu_khong_xac_dinh = []
            for muc in muc_thieu:
                dg = don_gia_thap_nhat.get(muc["ten"])
                if dg is None and args.du_toan:
                    dg = Decimal(str(args.du_toan))
                # KHÔNG mặc định số lượng = 1 khi HSMT không cho số lượng rõ ràng — số tiền ra
                # từ một số lượng bịa sẽ trực tiếp thành giá hợp đồng, đúng thứ "không bịa số"
                # skill này cấm.
                sl = Decimal(str(muc["so_luong"])) if muc.get("so_luong") else None
                if dg is None or sl is None:
                    thieu_khong_xac_dinh.append(muc["ten"])
                    continue
                gia_tri_khoan4 += dg * sl

            hang_thang = next(r for r in bang_tong if r["Nhà thầu"] == nguoi_thang)
            if not thieu_khong_xac_dinh:
                hieu_chinh_khoan4 = gia_tri_khoan4 - hang_thang["Sai lệch thừa"]
                gia_de_nghi_khoan4 = (hang_thang["Giá dự thầu"] + hang_thang["Giá trị sửa lỗi"]
                                      + hieu_chinh_khoan4 - hang_thang["Thư giảm giá"])
                hang_thang["Giá đề nghị trúng thầu (dự kiến, khoản 4 Điều 31)"] = gia_de_nghi_khoan4

            bang_canh_bao.append({
                "Nhà thầu": nguoi_thang, "Loại": "Khoản 4 Điều 31",
                "Nội dung": (
                    f"{nguoi_thang} đang xếp hạng nhất và có chào thiếu không tự có đơn giá "
                    "riêng. Theo khoản 4 Điều 31 NĐ 214/2025, giá đề nghị trúng thầu (giá đưa vào "
                    "hợp đồng) cho phần chào thiếu này phải tính bằng đơn giá THẤP NHẤT trong các "
                    "HSDT vượt kỹ thuật — ngược chiều với đơn giá CAO NHẤT đã dùng để xếp hạng ở "
                    "cột 'Giá đánh giá (G)'. Xem cột 'Giá đề nghị trúng thầu (dự kiến, khoản 4 "
                    "Điều 31)'. Nếu nhà thầu này KHÔNG trúng thầu sau cùng, cột này không áp dụng."
                    + (f" KHÔNG xác định được đơn giá và/hoặc số lượng cho: "
                       f"{', '.join(thieu_khong_xac_dinh)} — cột giá đề nghị trúng thầu dự kiến "
                       "để trống, cần --du-toan hoặc xác định thủ công." if thieu_khong_xac_dinh
                       else "")
                ),
            })

    tom_tat = {
        "so_nha_thau": len(bang_tong),
        "ky_thuat": {
            "vuot": sorted(vuot_kt), "khong_vuot": sorted(truot_kt),
            "chua_co_ket_luan": sorted(chua_ro),
            "nguon": args.ket_qua_ky_thuat or "CHƯA CÓ — đang tạm coi tất cả là vượt",
        },
        "so_chao_thieu": sum(1 for r in bang_doi_chieu if r["Tình trạng"] == "CHÀO THIẾU"),
        "so_chao_thua": sum(1 for r in bang_doi_chieu if r["Tình trạng"] == "CHÀO THỪA"),
        "so_loi_so_hoc": len(bang_sua_loi),
        "so_canh_bao": len(bang_canh_bao),
        "xep_hang_so_bo": [{"hang": r.get("Xếp hạng sơ bộ"), "nha_thau": r["Nhà thầu"],
                            "gia_danh_gia": r["Giá đánh giá (G)"]} for r in co_gia],
        "gia_de_nghi_trung_thau_du_kien_khoan_4": (
            {"nha_thau": co_gia[0]["Nhà thầu"], "gia_tri": gia_de_nghi_khoan4}
            if co_gia and gia_de_nghi_khoan4 is not None else None),
        "luu_y": LUU_Y,
        "can_cu_phap_ly": "NĐ 214/2025/NĐ-CP (KHÔNG dùng NĐ 63/2014 hay NĐ 24/2024 — đã hết hiệu lực)",
        "canh_bao_chung": canh_bao_chung,
    }

    if args.ra:
        ra = Path(args.ra).expanduser()
        ra.parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(ra, engine="openpyxl") as bo:
            # startrow=2 chừa dòng 1 cho lưu ý (ghi sau bằng openpyxl), header rơi vào dòng 3
            pd.DataFrame(bang_tong).to_excel(
                bo, sheet_name="Tong hop", index=False, startrow=2)
            for ten_sheet, du in (("Doi chieu danh muc", bang_doi_chieu),
                                  ("Sua loi", bang_sua_loi),
                                  ("Canh bao", bang_canh_bao),
                                  ("Can cu", bang_can_cu)):
                pd.DataFrame(du or [{"(không có)": ""}]).to_excel(
                    bo, sheet_name=ten_sheet[:31], index=False)
        _ghi_luu_y(ra, LUU_Y, len(bang_tong[0]) if bang_tong else 8)
        tom_tat["tep_ket_qua"] = str(ra)

    for c in canh_bao_chung:
        canh_bao(c)
    canh_bao(LUU_Y)

    in_json({"tom_tat": tom_tat, "tong_hop": bang_tong,
             "doi_chieu_danh_muc": bang_doi_chieu, "sua_loi": bang_sua_loi,
             "canh_bao": bang_canh_bao})
    return 0


if __name__ == "__main__":
    sys.exit(main())
