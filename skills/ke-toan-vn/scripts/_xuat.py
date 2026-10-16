"""Ghi workbook kết quả theo một khuôn chung: tiêu đề đậm, cố định dòng đầu, có bộ lọc."""

from __future__ import annotations

import os
from pathlib import Path


def sheet_du_lieu(wb, ten: str, tieu_de: list, dong: list):
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(ten)
    ws.append(tieu_de)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in dong:
        ws.append(r)
    ws.freeze_panes = "A2"
    if tieu_de:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(tieu_de))}{len(dong) + 1}"
    return ws


def dam_dong(ws, *so_dong: int) -> None:
    from openpyxl.styles import Font

    for r in so_dong:
        for c in ws[r]:
            c.font = Font(bold=True)


class BoGhi:
    """Gom mọi file kết quả của MỘT lần chạy và chỉ công bố khi mọi bước kiểm đã qua.

    Mỗi workbook được lưu vào file tạm cạnh đích (cùng ổ đĩa, nên đổi tên là nguyên tử). Chỉ khi
    `cong_bo()` mới đổi tên sang tên thật; `huy()` xoá file tạm. Nhờ vậy: dừng để hỏi kế toán,
    phát hiện bố cục đổi, hay lỗi ở file thứ hai của một lô — đều không để lại file kết quả ghi
    dở, và file kết quả cũ (nếu có) còn nguyên. `xem_truoc=True` thì không ghi gì cả.
    """

    def __init__(self, xem_truoc: bool = False):
        self.xem_truoc = xem_truoc
        self.tep: list[tuple[Path | None, Path]] = []

    def ghi(self, wb, ra: Path) -> None:
        wb.calculation.fullCalcOnLoad = True
        if self.xem_truoc:
            self.tep.append((None, ra))
            return
        ra.parent.mkdir(parents=True, exist_ok=True)
        tam = ra.with_name(f".~{os.getpid()}-{len(self.tep)}-{ra.name}")
        try:
            wb.save(tam)
        except BaseException:
            tam.unlink(missing_ok=True)
            raise
        self.tep.append((tam, ra))

    def doi_dich(self, cu: Path, moi: Path) -> None:
        self.tep = [(t, moi if ra == cu else ra) for t, ra in self.tep]

    def cong_bo(self) -> None:
        if self.xem_truoc:
            return
        hong = [str(ra) for _, ra in self.tep if ra.is_dir() or not os.access(ra.parent, os.W_OK)]
        if hong:
            raise OSError(f"Không ghi được vào: {hong} — chưa file kết quả nào được công bố")
        for tam, ra in self.tep:
            os.replace(tam, ra)
        self.tep = []

    def huy(self) -> None:
        for tam, _ in self.tep:
            if tam is not None:
                tam.unlink(missing_ok=True)
        self.tep = []


def luu(wb, ra: Path, bo_ghi: BoGhi | None = None) -> None:
    """Đặt tính lại khi mở: kế toán mở bằng Excel là thấy số, không cần recalc trước.

    Không có `bo_ghi` thì vẫn ghi nguyên tử (file tạm rồi đổi tên): máy tắt giữa chừng không để
    lại một file .xlsx hỏng thay cho file cũ.
    """
    if bo_ghi is not None:
        bo_ghi.ghi(wb, ra)
        return
    mot = BoGhi()
    try:
        mot.ghi(wb, ra)
        mot.cong_bo()
    finally:
        mot.huy()
