#!/usr/bin/env python3
"""Lọc bản kết xuất doanh thu của một khoa thành DLBC (dữ liệu báo cáo).

    loc_dlbc.py MAT-09.xlsx NT-09.xlsx ...            # ghi DLBC-Khoa-MAT-09.xlsx cạnh file gốc
    loc_dlbc.py YHCT-09.xlsx --ra-thu-muc ket_qua/

Mỗi file đầu vào sinh một file DLBC-Khoa-<tên file> gồm:
  - DLBC     : các dòng được giữ, đúng thứ tự cột của bản kết xuất — là đầu vào của
               bao_cao_khoa.py, mà công thức ở đó trỏ cột theo chữ cái.
  - LOẠI BỎ  : các dòng bị lọc, thêm cột cuối ghi quy tắc nào loại dòng đó.
  - TÓM TẮT  : nguồn, quy tắc áp dụng, và đối chiếu bằng công thức: số dòng và tổng
               SOTIENCT của gốc = DLBC + LOẠI BỎ.

Quy tắc (kế toán bệnh viện chốt tháng 9/2026). Một dòng bị loại nếu dính BẤT KỲ quy tắc nào:
  1. TENQUYENSO là quyển HDDTSOFTDREAM_*.
  2. TENNHOMBHYT thuộc CDHA, TDCN / TDCN / Khác / Máu / Xét nghiệm — trừ các chỉ định
     Đường máu mao mạch (kể cả [Ngoại trú]) và Định nhóm máu tại giường [...].
  3. TENLOAIVP thuộc Chăm sóc sức khỏe / Y học cổ truyền / Test tâm lý /
     Thủ thuật Tai Mũi Họng / Vật lý trị liệu - Phục hồi chức năng.
  4. SOTIENCT âm.
Riêng khoa YHCT (tên file bắt đầu bằng YHCT, hoặc cờ --yhct), quy tắc 3 không loại
Y học cổ truyền và không loại chỉ định Điều trị bằng tia hồng ngoại.

So khớp tên bỏ qua khoảng trắng thừa và hoa/thường: bản kết xuất ghi "Xét nghiệm " có dấu
cách cuối, và cùng một loại có lúc viết "vitamin", lúc "Vitamin".

Đã kiểm trên dữ liệu tháng 8/2026: áp các quy tắc lên sheet GỐC rồi so với DLBC kế toán tự
lọc. NTK khớp từng dòng. Phần lệch còn lại đều giải thích được:
  - MẮT: tháng 8 chưa loại HDDTSOFTDREAM (123 dòng) — quy tắc 1 mới có từ tháng 9.
  - YHCT: tháng 8 loại tia hồng ngoại (78 dòng) — quy tắc riêng YHCT mới có từ tháng 9; và
    kế toán thêm tay 183 dòng Y học cổ truyền không có trong GỐC.
  - NGOẠI TK 1, NỘI TIẾT 2, NỘI TK 10 dòng kế toán loại thêm ngoài quy tắc (giải phẫu bệnh,
    gây mê chụp MRI, chi phí vận chuyển, vài lượt khám) — báo cho kế toán, không tự loại.
"""

from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _chung import can_thu_vien, in_json  # noqa: E402

TIEN_TO_RA = "DLBC-Khoa-"
SHEET_GIU = "DLBC"
SHEET_BO = "LOẠI BỎ"
SHEET_TOM_TAT = "TÓM TẮT"
COT_LY_DO = "LÝ DO LOẠI"

QUYEN_SO_BO = "HDDTSOFTDREAM"
NHOM_BHYT_BO = ("CDHA, TDCN", "TDCN", "Khác", "Máu", "Xét nghiệm")
# So theo tiền tố: tên thật dài hơn, ví dụ "Định nhóm máu tại giường [của túi máu toàn phần...]".
CHI_DINH_GIU = ("Đường máu mao mạch", "Định nhóm máu tại giường")
LOAI_VP_BO = (
    "Chăm sóc sức khỏe", "Y học cổ truyền", "Test tâm lý",
    "Thủ thuật Tai Mũi Họng", "Vật lý trị liệu - Phục hồi chức năng",
)
YHCT_LOAI_VP_GIU = ("Y học cổ truyền",)
YHCT_CHI_DINH_GIU = ("Điều trị bằng tia hồng ngoại",)

COT_CAN = ("TENQUYENSO", "TENNHOMBHYT", "TENLOAIVP", "TENCHIDINH", "SOTIENCT")


def chuan(v) -> str:
    """Chuẩn hoá tên để so khớp: NFC, gộp khoảng trắng, bỏ phân biệt hoa/thường."""
    if v is None:
        return ""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(v))).strip().casefold()


def _tap(ten) -> tuple[str, ...]:
    return tuple(chuan(x) for x in ten)


def ly_do_loai(quyen_so, nhom_bhyt, loai_vp, chi_dinh, so_tien, yhct: bool) -> list[str]:
    """Danh sách quy tắc loại dòng này; rỗng nghĩa là giữ."""
    ly_do = []
    if chuan(quyen_so).startswith(chuan(QUYEN_SO_BO)):
        ly_do.append(f"Quy tắc 1: TENQUYENSO {quyen_so}")
    cd = chuan(chi_dinh)
    if chuan(nhom_bhyt) in _tap(NHOM_BHYT_BO) and not cd.startswith(_tap(CHI_DINH_GIU)):
        ly_do.append(f"Quy tắc 2: TENNHOMBHYT {str(nhom_bhyt).strip()}")
    lvp = chuan(loai_vp)
    mien_yhct = yhct and (lvp in _tap(YHCT_LOAI_VP_GIU) or cd.startswith(_tap(YHCT_CHI_DINH_GIU)))
    if lvp in _tap(LOAI_VP_BO) and not mien_yhct:
        ly_do.append(f"Quy tắc 3: TENLOAIVP {str(loai_vp).strip()}")
    if isinstance(so_tien, (int, float)) and so_tien < 0:
        ly_do.append("Quy tắc 4: SOTIENCT âm")
    return ly_do


def doc_nguon(duong_dan: Path, sheet: str | None):
    import openpyxl

    wb = openpyxl.load_workbook(duong_dan, read_only=True, data_only=True)
    ten_sheet = sheet or wb.sheetnames[0]
    if ten_sheet not in wb.sheetnames:
        raise SystemExit(f"Không có sheet {ten_sheet!r} trong {duong_dan.name}")
    it = wb[ten_sheet].iter_rows(values_only=True)
    tieu_de = list(next(it))
    # Cắt cột rỗng ở cuối tiêu đề: read_only trả về tới max_column của sheet.
    while tieu_de and tieu_de[-1] is None:
        tieu_de.pop()
    n = len(tieu_de)
    dong = [list(r[:n]) for r in it if not all(v is None for v in r)]
    wb.close()
    chi_muc = {}
    for i, h in enumerate(tieu_de):
        if h and h not in chi_muc:
            chi_muc[h] = i
    thieu = [c for c in COT_CAN if c not in chi_muc]
    if thieu:
        raise SystemExit(f"{duong_dan.name}: sheet {ten_sheet!r} thiếu cột {', '.join(thieu)}")
    return ten_sheet, tieu_de, chi_muc, dong


def ghi_ket_qua(ra: Path, nguon: Path, ten_sheet: str, tieu_de, chi_muc, giu, bo, yhct: bool):
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    dam = Font(bold=True)

    def sheet_du_lieu(ten, dau, dong):
        ws = wb.create_sheet(ten)
        ws.append(dau)
        for c in ws[1]:
            c.font = dam
        for r in dong:
            ws.append(r)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(dau))}{len(dong) + 1}"
        return ws

    sheet_du_lieu(SHEET_GIU, tieu_de, giu)
    ws_bo = sheet_du_lieu(SHEET_BO, tieu_de + [COT_LY_DO], bo)
    ws_bo.column_dimensions[get_column_letter(len(tieu_de) + 1)].width = 60

    # Đối chiếu bằng công thức để kế toán bấm vào ô thấy nó đếm/cộng từ đâu. Riêng số liệu
    # của file gốc là giá trị, vì file gốc không nằm trong workbook này.
    st = get_column_letter(chi_muc["SOTIENCT"] + 1)
    ld = get_column_letter(len(tieu_de) + 1)
    tong_goc = sum(r[chi_muc["SOTIENCT"]] or 0 for r in giu + bo
                   if isinstance(r[chi_muc["SOTIENCT"]], (int, float)))
    g, b = f"'{SHEET_GIU}'", f"'{SHEET_BO}'"
    ws = wb.create_sheet(SHEET_TOM_TAT)
    dong = [
        ["Nguồn", f"{nguon.name} / sheet {ten_sheet}"],
        ["Ngày lập", datetime.now().strftime("%d/%m/%Y %H:%M")],
        ["Quy tắc riêng YHCT", "Có" if yhct else "Không"],
        [],
        ["", "Số dòng", "Tổng SOTIENCT"],
        ["File gốc", len(giu) + len(bo), tong_goc],
        [f"Giữ lại ({SHEET_GIU})", f"=COUNTA({g}!A:A)-1", f"=SUM({g}!{st}:{st})"],
        [f"Loại bỏ ({SHEET_BO})", f"=COUNTA({b}!A:A)-1", f"=SUM({b}!{st}:{st})"],
        ["Gốc − giữ − loại", "=B6-B7-B8", "=C6-C7-C8"],
        ["Kết luận", '=IF(AND(B9=0,ABS(C9)<0.5),"Khớp","LỆCH")'],
        [],
        ["Số dòng dính từng quy tắc (một dòng có thể dính nhiều quy tắc)"],
    ]
    for r in dong:
        ws.append(r)
    for i in range(1, 5):
        ws.append([f"Quy tắc {i}", f'=COUNTIF({b}!{ld}:{ld},"*Quy tắc {i}:*")'])
    for r in (5, 12):
        for c in ws[r]:
            c.font = dam
    for r in range(6, 10):
        ws.cell(r, 3).number_format = "#,##0"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 20

    wb.calculation.fullCalcOnLoad = True
    ra.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ra)


def loc_mot_file(nguon: Path, ra_thu_muc: Path | None, sheet: str | None, yhct: bool | None):
    yhct = nguon.name.upper().startswith("YHCT") if yhct is None else yhct
    ten_sheet, tieu_de, chi_muc, dong = doc_nguon(nguon, sheet)
    giu, bo, dem = [], [], {}
    lay = [chi_muc[c] for c in ("TENQUYENSO", "TENNHOMBHYT", "TENLOAIVP", "TENCHIDINH", "SOTIENCT")]
    for r in dong:
        ly_do = ly_do_loai(*(r[i] for i in lay), yhct=yhct)
        if not ly_do:
            giu.append(r)
            continue
        bo.append(r + ["; ".join(ly_do)])
        for x in ly_do:
            dem[x] = dem.get(x, 0) + 1

    ra = (ra_thu_muc or nguon.parent) / f"{TIEN_TO_RA}{nguon.stem}.xlsx"
    ghi_ket_qua(ra, nguon, ten_sheet, tieu_de, chi_muc, giu, bo, yhct)
    st = chi_muc["SOTIENCT"]
    tong = lambda ds: round(sum(r[st] for r in ds if isinstance(r[st], (int, float))), 2)  # noqa: E731
    return {
        "nguon": str(nguon),
        "ra": str(ra),
        "quy_tac_yhct": yhct,
        "so_dong_goc": len(dong),
        "so_dong_giu": len(giu),
        "so_dong_bo": len(bo),
        "tong_sotienct_giu": tong(giu),
        "tong_sotienct_bo": tong([r[:-1] for r in bo]),
        "ly_do_bo": dict(sorted(dem.items(), key=lambda kv: -kv[1])),
    }


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("tep", nargs="+", help="Bản kết xuất doanh thu của khoa (.xlsx)")
    p.add_argument("--sheet", help="Sheet dữ liệu; mặc định là sheet đầu tiên")
    p.add_argument("--ra-thu-muc", help="Thư mục ghi kết quả; mặc định cạnh file gốc")
    nhom = p.add_mutually_exclusive_group()
    nhom.add_argument("--yhct", dest="yhct", action="store_true", default=None,
                      help="Áp quy tắc riêng YHCT dù tên file không bắt đầu bằng YHCT")
    nhom.add_argument("--khong-yhct", dest="yhct", action="store_false")
    a = p.parse_args()
    can_thu_vien("openpyxl")

    ra_thu_muc = Path(a.ra_thu_muc) if a.ra_thu_muc else None
    in_json([loc_mot_file(Path(t), ra_thu_muc, a.sheet, a.yhct) for t in a.tep])
    return 0


if __name__ == "__main__":
    sys.exit(main())
