#!/usr/bin/env python3
"""Kiểm thử hồi quy: chạy từng script trên dữ liệu mẫu và đối chiếu con số đã biết.

    python evals/kiem_thu.py

Chạy được ở cả máy cá nhân lẫn CI. Không cần LibreOffice hay tesseract — chỉ kiểm phần
logic nghiệp vụ, là phần dễ hỏng âm thầm nhất khi sửa code.

Mỗi phép kiểm gắn với một cái bẫy cố ý cài trong dữ liệu mẫu, nên khi nó đỏ thì biết ngay
nghiệp vụ nào vừa vỡ, chứ không phải "có gì đó sai".
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

GOC = Path(__file__).resolve().parent.parent
SCRIPTS = GOC / "skills" / "ke-toan-vn" / "scripts"
FIX = GOC / "evals" / "fixtures"
PY = sys.executable

loi: list[str] = []


def chay(*lenh: str) -> dict:
    kq = subprocess.run([PY, *lenh], capture_output=True, text=True)
    if kq.returncode != 0:
        raise SystemExit(f"Script lỗi: {' '.join(lenh)}\n{kq.stderr[-2000:]}")
    return json.loads(kq.stdout)


def _chuan(v):
    """Số tiền là Decimal, qua JSON thành chuỗi — so sánh phải chịu được cả hai kiểu."""
    if isinstance(v, str):
        try:
            return float(v)
        except ValueError:
            return v
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, list):
        return [_chuan(x) for x in v]
    return v


def kiem(ten: str, thuc_te, mong_doi) -> None:
    if _chuan(thuc_te) == _chuan(mong_doi):
        print(f"  ✅ {ten}")
    else:
        print(f"  ❌ {ten}: được {thuc_te!r}, mong đợi {mong_doi!r}")
        loi.append(ten)


print("Hóa đơn — XML + PDF, thuế suất 8% và KCT, MST sai checksum")
d = chay(str(SCRIPTS / "doc_hoa_don.py"), str(FIX / "hoa-don"))
t = d["tom_tat"]
kiem("đọc đủ 7 hóa đơn (5 XML + 2 PDF)", (t["so_hoa_don"], t["tu_xml"], t["tu_pdf"]), (7, 5, 2))
kiem("tổng tiền hàng", t["tong_tien_hang"], 145180000)
kiem("tổng tiền thuế", t["tong_tien_thue"], 12248000)
kiem("không có hóa đơn trùng số", t["hoa_don_trung"], [])
ts = {h["so_hd"]: h["thue_suat"] for h in d["chi_tiet"]}
kiem("giữ thuế suất 8%, không ép về 10%", ts["00000104"], 8.0)
kiem("giữ KCT, không quy thành 0%", ts["00000105"], "KCT")
mst_xau = [h for h in d["chi_tiet"] if "0100686200" in (h.get("canh_bao") or "")]
kiem("cảnh báo MST 0100686200 sai chữ số kiểm tra", len(mst_xau), 1)

print("\nMã số thuế")
d = chay(str(SCRIPTS / "kiem_tra_mst.py"), "0100109106", "0101243150", "0100686200")
kiem("2 hợp lệ, 1 nghi ngờ", d["tom_tat"], {"hop_le": 2, "nghi_ngo": 1})

print("\nĐối chiếu — thiếu, lệch tiền, ghi trùng")
d = chay(str(SCRIPTS / "doi_chieu.py"),
         str(FIX / "doi-chieu/bang_ke_ban_ra.csv"), str(FIX / "doi-chieu/so_chi_tiet_511.csv"),
         "--khoa-a", "so_hd", "--khoa-b", "so_chung_tu",
         "--tien-a", "tien_hang", "--tien-b", "phat_sinh_co")
t = d["tom_tat"]
kiem("1 hóa đơn có trong bảng kê mà thiếu ở sổ", t["so_chi_co_a"], 1)
kiem("không báo giả chiều ngược lại", t["so_chi_co_b"], 0)
kiem("phát hiện 1 dòng ghi trùng trong sổ", t["so_trung_lap_b"], 1)
kiem("hóa đơn thiếu đúng là 00000205", d["chi_co_a"][0]["khoa"], "205")
lech = {r["khoa"]: r["chenh_lech"] for r in d["lech_tien"]}
kiem("00000207 lệch đúng 900.000", lech.get("207"), "900000")
kiem("chênh lệch tổng giải thích được hết",
     "CÒN" not in t["doi_chieu_tong"]["ket_luan"], True)

print("\nHSDT — trích xuất và bắt lỗi số học")
d = chay(str(SCRIPTS / "doc_hsdt.py"), str(FIX / "goi-thau/hsdt"))
kiem("đọc đủ 4 bộ HSDT", d["tom_tat"]["so_nha_thau"], 4)
theo_ma = {nt["nha_thau"].split(" - ")[0]: nt for nt in d["nha_thau"]}
kiem("nhà thầu A chào thiếu 1 mặt hàng (7/8)", len(theo_ma["Nha thau A"]["danh_muc"]), 7)
kiem("bắt lỗi số học của nhà thầu B",
     sum(1 for x in theo_ma["Nha thau B"]["danh_muc"] if x.get("loi_so_hoc")), 1)
kiem("đọc thư giảm giá của nhà thầu C",
     theo_ma["Nha thau C"]["gia_giam"]["gia_tri"], "35448180")
kiem("đọc hiệu lực bảo đảm dự thầu của D",
     theo_ma["Nha thau D"]["bao_lanh_hieu_luc_den"]["gia_tri"], "2026-10-15")
kiem("mọi trường trích được đều có nguồn trang",
     all(theo_ma["Nha thau A"][k]["nguon_trang"] for k in ("gia_du_thau", "mst")), True)

print("\nChấm kỹ thuật — phân biệt sai lệch khách quan với việc cần người xem")
tmp = Path("/tmp/kt_kiem_thu")
tmp.mkdir(exist_ok=True)
d = chay(str(SCRIPTS / "cham_ky_thuat.py"), str(FIX / "goi-thau/hsdt"),
         "--yeu-cau", str(FIX / "goi-thau/yeu_cau_ky_thuat.csv"),
         "--ra", str(tmp / "phieu.xlsx"))
theo = d["tom_tat"]["theo_nha_thau"]
lay = lambda ma, k: next(v[k] for n, v in theo.items() if n.startswith(f"Nha thau {ma}"))  # noqa: E731
kiem("C: RAM 8GB < 16GB bị bắt bằng số học", lay("C", "thap_hon"), 1)
kiem("B: thiếu hẳn tiêu chí bảo hành", lay("B", "thieu_du_lieu"), 1)
kiem("D: xuất xứ khác loại → để người xem, không tự loại", lay("D", "thap_hon"), 0)
kiem("không tự điền kết luận cho tiêu chí nào",
     all(h["ket_luan"] == "" for h in d["chi_tiet"]), True)

print("\nGiá đánh giá — thứ tự tính và ảnh hưởng của bước kỹ thuật")
d = chay(str(SCRIPTS / "doc_hsdt.py"), str(FIX / "goi-thau/hsdt"), "--ra", str(tmp / "h.json"))
d = chay(str(SCRIPTS / "so_sanh_thau.py"), str(tmp / "h.json"),
         "--danh-muc", str(FIX / "goi-thau/danh_muc_hsmt.csv"))
g = {r["Nhà thầu"].split(" - ")[0]: r for r in d["tong_hop"]}
kiem("sửa lỗi của B đúng -650.000", g["Nha thau B"]["Giá trị sửa lỗi"], "-650000")
kiem("giảm giá của C trừ SAU sửa lỗi và hiệu chỉnh",
     g["Nha thau C"]["Giá đánh giá (G)"], "1146157820")
kiem("chưa có kết quả kỹ thuật thì phải cảnh báo",
     any("kỹ thuật" in c for c in d["tom_tat"]["canh_bao_chung"]), True)
kiem("hiệu chỉnh dùng đơn giá CAO NHẤT (của C), không phải thấp nhất",
     next(r["Thành tiền"] for r in d["doi_chieu_danh_muc"] if r["Tình trạng"] == "CHÀO THIẾU"),
     "73542000")

# Tổ chuyên gia kết luận C không đạt -> C rời nhóm vượt kỹ thuật -> đơn giá hiệu chỉnh đổi
import openpyxl  # noqa: E402

wb = openpyxl.load_workbook(tmp / "phieu.xlsx")
ws = wb["Phieu cham ky thuat"]
hdr = {str(c.value).strip(): int(c.column or 0) for c in ws[3]}
for r in range(4, ws.max_row + 1):
    nt = str(ws.cell(row=r, column=hdr["nha_thau"]).value or "")
    ws.cell(row=r, column=hdr["ket_luan"]).value = (
        "Không đạt" if nt.startswith("Nha thau C") else "Đạt")
wb.save(tmp / "phieu.xlsx")

d = chay(str(SCRIPTS / "so_sanh_thau.py"), str(tmp / "h.json"),
         "--danh-muc", str(FIX / "goi-thau/danh_muc_hsmt.csv"),
         "--ket-qua-ky-thuat", str(tmp / "phieu.xlsx"))
kiem("C bị loại khỏi nhóm vượt kỹ thuật",
     [x.split(" - ")[0] for x in d["tom_tat"]["ky_thuat"]["khong_vuot"]], ["Nha thau C"])
kiem("đơn giá hiệu chỉnh tụt xuống mức của D",
     next(r["Thành tiền"] for r in d["doi_chieu_danh_muc"] if r["Tình trạng"] == "CHÀO THIẾU"),
     "70686000")
kiem("C không còn trong bảng xếp hạng",
     [x["nha_thau"].split(" - ")[0] for x in d["tom_tat"]["xep_hang_so_bo"]],
     ["Nha thau B", "Nha thau D", "Nha thau A"])

print("\nKhoản 4 Điều 31 — đơn giá THẤP NHẤT cho nhà thầu xếp hạng nhất (ngược chiều khoản 2)")
# Dữ liệu dựng tay (không qua doc_hsdt.py): NT Thang chào thiếu 'Muc X' (10 cái), nhưng giá dự
# thầu đủ thấp để vẫn xếp hạng nhất SAU KHI bị cộng hiệu chỉnh bằng đơn giá CAO NHẤT (300, của
# NT Ba) theo khoản 2. Khoản 4 phải dùng đơn giá THẤP NHẤT (200, của NT Hai) khi tính giá đề
# nghị trúng thầu — một con số khác, thấp hơn, cho đúng nhà thầu đó.
d = chay(str(SCRIPTS / "so_sanh_thau.py"), str(FIX / "goi-thau/khoan4_hsdt.json"),
         "--danh-muc", str(FIX / "goi-thau/khoan4_danh_muc.csv"))
g = {r["Nhà thầu"]: r for r in d["tong_hop"]}
kiem("NT Thang xếp hạng nhất dù chào thiếu", g["NT Thang"]["Xếp hạng sơ bộ"], 1)
kiem("giá đánh giá (G) dùng đơn giá CAO NHẤT (300×10=3000 cộng vào 10000)",
     g["NT Thang"]["Giá đánh giá (G)"], "13000")
kiem("giá đề nghị trúng thầu dự kiến dùng đơn giá THẤP NHẤT (200×10=2000 cộng vào 10000)",
     g["NT Thang"]["Giá đề nghị trúng thầu (dự kiến, khoản 4 Điều 31)"], "12000")
kiem("nhà thầu không xếp hạng nhất thì không có giá đề nghị trúng thầu dự kiến",
     (g["NT Hai"]["Giá đề nghị trúng thầu (dự kiến, khoản 4 Điều 31)"],
      g["NT Ba"]["Giá đề nghị trúng thầu (dự kiến, khoản 4 Điều 31)"]), (None, None))
kiem("NT Tu có chào 'Muc X' nhưng thiếu đơn giá riêng -> cảnh báo, không bịa số",
     any(c["Nhà thầu"] == "NT Tu" and c["Loại"] == "Thiếu đơn giá dòng đã chào"
         for c in d["canh_bao"]), True)
kiem("đơn giá thiếu của NT Tu không lẫn vào đơn giá cao/thấp nhất dùng chung",
     (g["NT Thang"]["Giá đánh giá (G)"],
      g["NT Thang"]["Giá đề nghị trúng thầu (dự kiến, khoản 4 Điều 31)"]), ("13000", "12000"))

print("\nBộ nhớ — hạn dùng và bản lưu trữ")
bn = str(SCRIPTS / "bo_nho.py")
import os  # noqa: E402

os.environ["KE_TOAN_VN_HOME"] = str(tmp / "bo-nho")
subprocess.run([PY, bn, "ghi", "--loai", "phap-ly", "--chu-de", "Thu nghiem",
                "--nguon", "https://vanban.chinhphu.vn/x", "--noi-dung", "Mức cũ 500 triệu."],
               capture_output=True, check=True)
subprocess.run([PY, bn, "ghi", "--loai", "phap-ly", "--chu-de", "Thu nghiem",
                "--nguon", "https://vanban.chinhphu.vn/y", "--noi-dung", "Mức mới 1 tỷ."],
               capture_output=True, check=True)
d = chay(bn, "tra", "thu nghiem")
kiem("tra cứu chỉ trả bản hiện hành, bỏ qua bản lưu trữ .cu-", d["so_ghi_chu"], 1)
kiem("nội dung là bản mới", "1 tỷ" in d["ghi_chu"][0]["noi_dung"], True)
kiem("ghi chú còn hạn thì không cần tra web", d["can_tra_web"], False)

print()
if loi:
    print(f"❌ {len(loi)} phép kiểm thất bại: {', '.join(loi)}")
    sys.exit(1)
print("✅ Tất cả phép kiểm đều đạt")
